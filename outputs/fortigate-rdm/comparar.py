"""Comparação pré/pós e relatório Excel. Nenhum acesso de rede neste módulo."""
import json
from pathlib import Path
import re

CATEGORIES = ['BGP', 'VPN', 'Rotas', 'Interfaces', 'Interfaces_cfg', 'ARP', 'DHCP',
              'DHCP_cfg', 'HA', 'SDWAN', 'Sistema', 'Recursos']


def compact(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list)) else str(value)


def compare(before, after):
    if before.get('schema') != 1 or after.get('schema') != 1:
        raise ValueError('Formato de coleta incompatível')
    if before.get('rdm') != after.get('rdm'):
        raise ValueError('As coletas pertencem a RDMs diferentes')
    if before.get('phase') != 'pre' or after.get('phase') != 'pos':
        raise ValueError('Selecione uma coleta pre e outra pos, nessa ordem')
    if after['started'] <= before['started']:
        raise ValueError('A coleta pós precisa ser posterior à coleta pré')
    alerts, coverage, hosts = [], [], []
    def add(device, area, level, event, key, old, new, action):
        alerts.append({'Equipamento': device, 'Área': area, 'Prioridade': level, 'Evento': event,
                       'Objeto': key, 'Antes': compact(old), 'Depois': compact(new), 'Verificação sugerida': action})
    for device in sorted(set(before['devices']) | set(after['devices'])):
        pre = before['devices'].get(device)
        post = after['devices'].get(device)
        if pre is None or post is None:
            add(device, 'Coleta', 'Alta', 'Equipamento ausente em uma coleta', device, bool(pre), bool(post), 'Conferir inventário e acesso SSH.')
            continue
        if (pre['host'], pre['port']) != (post['host'], post['port']):
            add(device, 'Coleta', 'Alta', 'Alvo alterado; comparação bloqueada', device, pre['host'], post['host'], 'Usar o mesmo endereço do cluster nas duas coletas.')
            continue
        for area in CATEGORIES:
            a = pre['data'].get(area, {'status': 'missing'})
            b = post['data'].get(area, {'status': 'missing'})
            valid = a['status'] == b['status'] == 'ok'
            coverage.append({'Equipamento': device, 'Área': area, 'Pré': a['status'], 'Pós': b['status'],
                             'Comparável': 'Sim' if valid else 'Não', 'Registros pré': len(a.get('rows', [])),
                             'Registros pós': len(b.get('rows', [])), 'Erro pré': a.get('error', ''), 'Erro pós': b.get('error', '')})
            if not valid:
                add(device, area, 'Alta', 'Comparação inconclusiva', '', a.get('error') or a['status'], b.get('error') or b['status'], 'Conferir logs e repetir a coleta; não inferir remoções.')
                continue
            aa, bb = a['rows'], b['rows']
            if area in ('ARP', 'DHCP'):
                cfg_a = {r['id']: r for r in pre['data'].get('Interfaces_cfg', {}).get('rows', [])}
                cfg_b = {r['id']: r for r in post['data'].get('Interfaces_cfg', {}).get('rows', [])}
                by_ip_a, by_ip_b = {}, {}
                for source, dest in ((aa, by_ip_a), (bb, by_ip_b)):
                    for row in source:
                        dest.setdefault(row['ip'], set()).add((row['mac'], row.get('interface', '')))
                for ip in sorted(set(by_ip_a) | set(by_ip_b)):
                    old, new = by_ip_a.get(ip, set()), by_ip_b.get(ip, set())
                    event = 'Mantido' if old == new else 'Não observado no pós' if not new else 'Novo no pós' if not old else 'MAC ou interface alterado'
                    old_vlan = sorted({cfg_a.get(i, {}).get('vlanid', 'não informado') for _, i in old})
                    new_vlan = sorted({cfg_b.get(i, {}).get('vlanid', 'não informado') for _, i in new})
                    if old == new and old_vlan != new_vlan:
                        event = 'VLAN da interface alterada'
                    hosts.append({'Equipamento': device, 'Fonte': area, 'IP': ip, 'Resultado': event,
                                  'MAC/interface pré': compact(sorted(old)), 'MAC/interface pós': compact(sorted(new)),
                                  'VLAN pré': ', '.join(old_vlan), 'VLAN pós': ', '.join(new_vlan)})
                    if event not in ('Mantido', 'Novo no pós'):
                        add(device, area, 'Revisar', event, ip, sorted(old), sorted(new),
                            'Conferir atividade, expiração ARP/lease, acesso e switch de origem; ausência não comprova queda.')
                mac_a, mac_b = {}, {}
                for source, dest in ((aa, mac_a), (bb, mac_b)):
                    for row in source:
                        dest.setdefault(row['mac'], set()).add((row['ip'], row.get('interface', '')))
                for mac in mac_a.keys() & mac_b.keys():
                    if mac_a[mac] != mac_b[mac]:
                        add(device, area, 'Revisar', 'Mesmo MAC observado com IP/interface diferente', mac,
                            sorted(mac_a[mac]), sorted(mac_b[mac]), 'Validar mudança esperada de endereço ou VLAN; MAC não é identidade garantida.')
                continue
            # Multiple ECMP paths remain one route object; uptime/counters are intentionally excluded.
            def index(rows):
                out = {}
                for row in rows:
                    stable = {k: v for k, v in row.items() if k not in ('uptime', 'rx_errors', 'tx_errors')}
                    out.setdefault(row['id'], []).append(stable)
                return {k: sorted(v, key=compact) for k, v in out.items()}
            ai, bi = index(aa), index(bb)
            for key in sorted(set(ai) | set(bi)):
                old, new = ai.get(key, []), bi.get(key, [])
                if old == new:
                    # Ongoing bad conditions also need attention even if they pre-date the RDM.
                    if area == 'HA' and any(r.get('value') == 'out-of-sync' for r in new):
                        add(device, area, 'Alta', 'HA fora de sincronismo (já existente)', key, old, new, 'Verificar sincronismo do cluster.')
                    continue
                event = 'Removido/não observado' if not new else 'Adicionado' if not old else 'Alterado'
                level = 'Revisar'
                if area == 'BGP' and old and old[0]['state'] == 'Established':
                    if not new or new[0]['state'] != 'Established':
                        event, level = 'Sessão BGP deixou de estar estabelecida', 'Alta'
                    elif new[0]['prefixes'] < old[0]['prefixes']:
                        event, level = 'Redução de prefixos BGP', 'Alta'
                if area == 'VPN' and old and old[0]['up'] > 0 and (not new or new[0]['up'] < old[0]['up']):
                    event, level = 'Redução/perda de seletores VPN ativos', 'Alta'
                if area == 'Interfaces' and old and old[0]['link'] == 'up' and (not new or new[0]['link'] != 'up'):
                    event, level = 'Interface perdeu link', 'Alta'
                if area == 'Rotas' and not new:
                    level = 'Alta'
                if area == 'HA' and (not new or any(r.get('value') == 'out-of-sync' for r in new)):
                    level = 'Alta'
                if area == 'HA' and key == 'HA Health Status' and new and new[0].get('value', '').upper() != 'OK':
                    event, level = 'Estado de saúde HA requer atenção', 'Alta'
                if area == 'SDWAN' and old and old[0]['state'] == 'alive' and (not new or new[0]['state'] != 'alive'):
                    event, level = 'Membro SD-WAN deixou de responder ao health-check', 'Alta'
                if area in ('Sistema', 'Recursos'):
                    level = 'Informação'
                add(device, area, level, event, key, old, new, 'Conferir se a alteração estava prevista na RDM e validar o serviço afetado.')
    alerts.sort(key=lambda r: ({'Alta': 0, 'Revisar': 1, 'Informação': 2}[r['Prioridade']], r['Equipamento'], r['Área']))
    return {'Alertas': alerts, 'Cobertura': coverage, 'Hosts': hosts}


def export_excel(before, after, report, output, pre_path=None, post_path=None):
    """Default output uses the user's template, extending rows without a fixed capacity."""
    from copy import copy
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.workbook.properties import CalcProperties
    template = Path(__file__).with_name('modelo_comparativo.xlsx')
    output = Path(output)
    if output.resolve() == template.resolve() or output.exists():
        raise ValueError('Destino já existe ou é o modelo; escolha outro nome')
    wb = load_workbook(template)
    if not {'Comparativo', 'Evidencias'} <= set(wb.sheetnames):
        raise ValueError('Modelo precisa conter Comparativo e Evidencias')
    ws, ev = wb['Comparativo'], wb['Evidencias']
    def put(sheet, row, column, value):
        cell = sheet.cell(row, column)
        cell.value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))[:32767] if value is not None else ''
        cell.data_type = 's'
    comp_style = [copy(ws.cell(15,c)._style) for c in range(1,13)]
    ev_style = [copy(ev.cell(6,c)._style) for c in range(1,9)]
    for sheet, start in ((ws,15),(ev,6)):
        for cells in sheet.iter_rows(min_row=start):
            for cell in cells:
                cell.value = None
    metadata = {'B4':before['rdm'], 'F4':', '.join(sorted(set(before['devices']) | set(after['devices']))),
                'B6':before['started'], 'D6':before.get('finished',''), 'F6':after['started'],
                'H6':after.get('finished',''), 'J6':'UTC', 'L6':'SSH / snapshots pré e pós',
                'B8':'Escopo das coletas selecionadas', 'D8':'Pendente'}
    for coord,value in metadata.items():
        cell=ws[coord]; put(ws,cell.row,cell.column,value)
    ws['A12']='Comparação preenchida pela automação. Complete resultado esperado, validação técnica e responsável.'
    ws['A13']='Igual não comprova funcionamento. Falhas de coleta ficam sem dados. Consulte Alertas, Cobertura e evidências.'
    ev['A3']='Evidências vinculadas aos itens do Comparativo. Saídas originais preservadas quando disponíveis.'
    ev['A4']='Linhas ampliadas automaticamente. Logs longos são divididos em partes; a fonte identifica o arquivo original.'
    references = {}
    evidence_rows = []
    for phase,snapshot,path in [('Pré',before,pre_path),('Pós',after,post_path)]:
        for device,entry in snapshot['devices'].items():
            for area,sample in entry['data'].items():
                source = str(path) if path else 'Snapshot em memória'
                raw, note = '', 'Status: ' + sample.get('status','missing')
                if path and sample.get('log'):
                    base = Path(path).resolve().parent
                    log = (base / sample['log']).resolve()
                    if log.is_relative_to(base) and log.is_file():
                        raw=log.read_text(encoding='utf-8')
                        source=str(log)
                    else:
                        note += '; arquivo de log não encontrado'
                if not raw:
                    raw=compact(sample.get('rows',[]))
                    note += '; dados estruturados do snapshot; não é saída CLI original'
                if sample.get('error'):
                    note += '; ' + sample['error']
                chunks=[raw[i:i+25000] for i in range(0,len(raw),25000)] or ['']
                ids=[]
                for part,chunk in enumerate(chunks,1):
                    eid=f'EV-{len(evidence_rows)+1:04}'
                    ids.append(eid)
                    evidence_rows.append([eid,phase,sample.get('at',snapshot['started']),device,
                                          sample.get('command',area),source,chunk,note+f'; parte {part}/{len(chunks)}'])
                references[phase,device,area]=', '.join(ids)
    for index,values in enumerate(evidence_rows,6):
        ev.row_dimensions[index].height=100
        for col,value in enumerate(values,1):
            put(ev,index,col,value)
            ev.cell(index,col)._style=copy(ev_style[col-1])
            ev.cell(index,col).alignment=Alignment(vertical='top',wrap_text=True)
    items=[]
    excluded={'age','uptime','rx_errors','tx_errors','details'}
    def index(sample):
        groups={}
        for record in sample.get('rows',[]):
            key=str(record.get('id',record.get('ip','sem identificador')))
            groups.setdefault(key,[]).append({k:v for k,v in record.items() if k not in excluded})
        return {key:sorted(rows,key=compact) for key,rows in groups.items()}
    for device in sorted(set(before['devices']) | set(after['devices'])):
        pre=before['devices'].get(device,{}); post=after['devices'].get(device,{})
        same_target=bool(pre and post and (pre['host'],pre['port']) == (post['host'],post['port']))
        for area in CATEGORIES:
            a=pre.get('data',{}).get(area,{}); b=post.get('data',{}).get(area,{})
            good_a=a.get('status')=='ok' and same_target
            good_b=b.get('status')=='ok' and same_target
            ai,bi=index(a) if good_a else {},index(b) if good_b else {}
            keys=sorted(set(ai)|set(bi)) or ['Visão geral']
            for key in keys:
                old=compact(ai[key]) if key in ai else 'Não observado nesta coleta' if good_a else ''
                new=compact(bi[key]) if key in bi else 'Não observado nesta coleta' if good_b else ''
                notes=[r['Evento']+': '+r['Verificação sugerida'] for r in report['Alertas']
                       if r['Equipamento']==device and r['Área']==area and (not r['Objeto'] or r['Objeto']==key)]
                if not good_a or not good_b:
                    notes.insert(0,'Comparação inconclusiva: coleta ausente/com falha ou alvo diferente. Consultar Cobertura.')
                if not ai and not bi and good_a and good_b:
                    notes.append('Sem registros nas duas coletas; confirmar aplicabilidade. Não comprova funcionamento.')
                if area in ('ARP','DHCP'):
                    notes.append('Ausência não comprova queda; expiração e atividade afetam ARP/leases. Conferir Hosts.')
                if area=='HA' and any(r.get('value','').lower()=='standalone' for r in a.get('rows',[])+b.get('rows',[])):
                    notes.append('Equipamento informou Standalone; não comprova funcionamento de cluster HA.')
                if len(old)>25000 or len(new)>25000:
                    raise ValueError('Objeto comparativo excede capacidade da célula; requer divisão antes de exportar')
                items.append([area,f'{device} / {key}','',old,new,'Pendente',
                              references.get(('Pré',device,area),''),references.get(('Pós',device,area),''),
                              '\n'.join(dict.fromkeys(notes)),''])
        for area,topic in [('Tráfego','Fluxos críticos'),('Políticas','Policy ID / ação / NAT'),
                           ('Segurança','Bloqueios e eventos'),('Aplicações','Testes de serviços'),
                           ('OSPF','Vizinhos OSPF'),('SD-WAN SLA','Latência, perda e jitter')]:
            items.append([area,f'{device} / {topic}','','','','Pendente','','','Sem coleta específica nesta automação; validação manual necessária.',''])
    for row,values in enumerate(items,15):
        ws.row_dimensions[row].height=90
        for col in range(1,13):
            ws.cell(row,col)._style=copy(comp_style[col-1])
            ws.cell(row,col).alignment=Alignment(vertical='top',wrap_text=True)
        put(ws,row,1,f'CMP-{row-14:04}')
        # B-F, H-L. G remains a native Excel formula from the template convention.
        for col,value in zip([2,3,4,5,6,8,9,10,11,12],values):
            put(ws,row,col,value)
        ws.cell(row,7).value=f'=IF(C{row}="","",IF(OR(E{row}="",F{row}=""),"Sem dados",IF(EXACT(E{row},F{row}),"Igual","Alterado")))'
    last=max(15,14+len(items)); end_ev=max(6,5+len(evidence_rows))
    ws['B10']=f'=COUNTA(C15:C{last})'
    ws['D10']=f'=COUNTIF(G15:G{last},"Alterado")'
    ws['F10']=f'=COUNTIFS(C15:C{last},"<>",H15:H{last},"Não conforme")'
    ws['H10']=f'=B10-COUNTIFS(C15:C{last},"<>",H15:H{last},"Conforme")-COUNTIFS(C15:C{last},"<>",H15:H{last},"Não conforme")-COUNTIFS(C15:C{last},"<>",H15:H{last},"Não aplicável")'
    for validation in ws.data_validations.dataValidation:
        if 'H15' in str(validation.sqref): validation.sqref=f'H15:H{last}'
    for validation in ev.data_validations.dataValidation:
        if 'B6' in str(validation.sqref): validation.sqref=f'B6:B{end_ev}'
    # Extend template conditional formats and tables to every generated item.
    for sheet,lastrow,start in [(ws,last,15),(ev,end_ev,6)]:
        for cf in list(sheet.conditional_formatting):
            rules=[copy(rule) for rule in sheet.conditional_formatting[cf]]
            ranges=[]
            for cell_range in cf.sqref.ranges:
                bounds=copy(cell_range)
                if bounds.min_row==start: bounds.max_row=lastrow
                ranges.append(str(bounds))
            del sheet.conditional_formatting[str(cf.sqref)]
            for rule in rules: sheet.conditional_formatting.add(' '.join(ranges),rule)
        for table in sheet.tables.values():
            from openpyxl.utils.cell import range_boundaries,get_column_letter
            lo,top,hi,_=range_boundaries(table.ref)
            table.ref=f'{get_column_letter(lo)}{top}:{get_column_letter(hi)}{lastrow}'
            from openpyxl.worksheet.filters import AutoFilter
            table.autoFilter=AutoFilter(ref=table.ref)
        sheet.auto_filter.ref=f'A{start-1}:{"L" if sheet==ws else "H"}{lastrow}'
        sheet.print_area=f'A1:{"L" if sheet==ws else "H"}{lastrow}'
    # Keep actionable host findings and completeness in the same output file.
    for name,rows in report.items():
        if name in wb.sheetnames: del wb[name]
        sheet=wb.create_sheet(name)
        headers=list(dict.fromkeys(k for r in rows for k in r)) or ['Informação']
        sheet.append(headers)
        if not rows: sheet.append(['Sem registros; consultar Cobertura.'])
        for row,record in enumerate(rows,2):
            for col,key in enumerate(headers,1): put(sheet,row,col,record.get(key,''))
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='17365D'); cell.font=Font(color='FFFFFF',bold=True)
        from openpyxl.utils import get_column_letter
        for col in range(1,len(headers)+1): sheet.column_dimensions[get_column_letter(col)].width=32
        sheet.freeze_panes='B2'; sheet.auto_filter.ref=sheet.dimensions
    wb.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
    output.parent.mkdir(parents=True,exist_ok=True)
    with output.open('xb') as handle: wb.save(handle)


def export_excel_legacy(before, after, report, output):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    output = Path(output)
    if output.exists():
        raise ValueError('Destino já existe; escolha outro nome para preservar o relatório anterior')
    sheets = {'Resumo': [
        {'Campo': 'RDM', 'Valor': before['rdm']}, {'Campo': 'Coleta pré UTC', 'Valor': before['started']},
        {'Campo': 'Coleta pós UTC', 'Valor': after['started']},
        {'Campo': 'Alertas de prioridade alta', 'Valor': sum(r['Prioridade'] == 'Alta' for r in report['Alertas'])},
        {'Campo': 'Áreas inconclusivas', 'Valor': sum(r['Comparável'] == 'Não' for r in report['Cobertura'])},
        {'Campo': 'Limite', 'Valor': 'Ausência em ARP/DHCP não comprova indisponibilidade. Revisar alertas e cobertura; não é aprovação automática da RDM.'}], **report}
    for phase, snapshot in [('Pré', before), ('Pós', after)]:
        for area in CATEGORIES:
            rows = []
            for name, device in snapshot['devices'].items():
                sample = device['data'].get(area, {})
                rows.extend({'Equipamento': name, 'Coleta UTC': sample.get('at', ''), **row} for row in sample.get('rows', []))
            sheets[phase + ' ' + area] = rows
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        headers = list(dict.fromkeys(k for r in rows for k in r)) or ['Informação']
        ws.append(headers)
        if not rows:
            ws.append(['Sem registros nesta aba; conferir Cobertura.'])
        for row in rows:
            ws.append([compact(row.get(k, '')) if isinstance(row.get(k), (dict, list)) else row.get(k, '') for k in headers])
        for cells in ws:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', cell.value)[:32767]
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(color='FFFFFF', bold=True)
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions
        for i, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = 55 if header in ('Antes', 'Depois', 'Verificação sugerida', 'Valor') else 26
        ws.sheet_view.showGridLines = False
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation also prevents a race from replacing another report.
    with output.open('xb') as handle:
        wb.save(handle)
