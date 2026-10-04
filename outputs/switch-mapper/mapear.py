"""Inventário SSH somente leitura para Cisco e Juniper. Python 3.10+."""
from __future__ import annotations

# Shared core is resolved for both source and portable distributions.
import sys as _sys
from pathlib import Path as _Path
for _parent in _Path(__file__).resolve().parents:
    if (_parent/'refresh_core').is_dir():
        if str(_parent) not in _sys.path:_sys.path.insert(0,str(_parent))
        break
else:raise ImportError('refresh_core ausente: use o pacote completo da Central')

import argparse
import csv
import getpass
import importlib
import json
import os
from pathlib import Path
import re
import shlex
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone


PLATFORMS = {"cisco_ios", "cisco_nxos", "juniper_junos"}


def vlan_set(value):
    value = value.strip().lower()
    if value == "all":
        return set(range(1, 4095))
    if value == "none":
        return set()
    result = set()
    for part in value.split(","):
        if "-" in part:
            a, b = map(int, part.split("-"))
            if not 1 <= a <= b <= 4094:
                raise ValueError("Intervalo VLAN inválido")
            result.update(range(a, b + 1))
        else:
            number = int(part)
            if not 1 <= number <= 4094:
                raise ValueError("VLAN inválida")
            result.add(number)
    return result


def compact(values):
    numbers = sorted(values)
    if not numbers:
        return "none"
    chunks = []
    start = end = numbers[0]
    for n in numbers[1:]:
        if n == end + 1:
            end = n
        else:
            chunks.append(str(start) if start == end else f"{start}-{end}")
            start = end = n
    chunks.append(str(start) if start == end else f"{start}-{end}")
    return ",".join(chunks)


def parse_cisco(config):
    interfaces, vlans, groups, notices = [], [], [], []
    current = None
    active_vlans = []
    allowed = {}
    for raw in config.splitlines():
        line = raw.strip()
        if line.startswith("interface "):
            name = line[10:]
            current = {"interface": name}
            interfaces.append(current)
            active_vlans = []
            continue
        if line.startswith("vlan ") and re.fullmatch(r"[\d, -]+", line[5:]):
            current = None
            active_vlans = [{"vlan_id": n} for n in sorted(vlan_set(line[5:].replace(" ", "")))]
            vlans.extend(active_vlans)
            continue
        if not raw.startswith((" ", "\t")):
            current, active_vlans = None, []
        if active_vlans and line.startswith("name "):
            for vlan in active_vlans:
                vlan["nome"] = line[5:]
        if current is None:
            continue
        fields = {"description ": "descricao", "switchport mode ": "modo_configurado",
                  "switchport access vlan ": "vlan_access", "switchport voice vlan ": "vlan_voz",
                  "switchport trunk native vlan ": "vlan_nativa", "speed ": "velocidade_configurada",
                  "duplex ": "duplex_configurado", "mtu ": "mtu"}
        for prefix, key in fields.items():
            if line.startswith(prefix):
                current[key] = line[len(prefix):]
        if line in ("shutdown", "no shutdown"):
            current["admin_configurado"] = "down" if line == "shutdown" else "up"
        if line == "no switchport":
            current["modo_configurado"] = "routed"
        if line.startswith(("ip address ", "ipv6 address ")):
            current["enderecos"] = (current.get("enderecos", "") + "; " + line).strip("; ")
        if line.startswith("switchport trunk allowed vlan "):
            expr = line.removeprefix("switchport trunk allowed vlan ")
            op, _, rest = expr.partition(" ")
            base = allowed.get(current["interface"], set(range(1, 4095)))
            if op == "add":
                base |= vlan_set(rest)
            elif op == "remove":
                base -= vlan_set(rest)
            elif op == "except":
                base = set(range(1, 4095)) - vlan_set(rest)
            else:
                base = vlan_set(expr)
            allowed[current["interface"]] = base
            current["vlans_trunk"] = compact(base)
        match = re.match(r"channel-group (\d+)(?: mode (\S+))?", line)
        if match:
            group, mode = match.groups()
            current["agregacao"] = "Port-channel" + group
            groups.append({"agregacao": "Port-channel" + group, "membro": current["interface"],
                           "modo": mode or "não explícito"})
    for row in interfaces:
        if row["interface"].startswith("range "):
            notices.append("interface range não expandida: " + row["interface"])
    return interfaces, vlans, groups, notices


def parse_junos(config):
    interfaces, vlans, groups, notices = {}, {}, [], []
    inactive = []
    lines = []
    for raw in config.splitlines():
        try:
            tokens = shlex.split(raw, comments=True)
        except ValueError:
            notices.append("Linha de configuração Junos não interpretada")
            continue
        if tokens and tokens[0] == "deactivate":
            inactive.append(tokens[1:])
        elif tokens and tokens[0] == "set":
            lines.append(tokens[1:])
    for t in lines:
        if any(t[:len(prefix)] == prefix for prefix in inactive):
            continue
        if len(t) >= 4 and t[0] == "vlans":
            vlan = vlans.setdefault(t[1], {"nome": t[1]})
            if t[2] in ("vlan-id", "vlan-id-list", "description", "l3-interface"):
                vlan[{"vlan-id": "vlan_id", "description": "descricao"}.get(t[2], t[2])] = " ".join(t[3:])
        if len(t) < 3 or t[0] != "interfaces":
            continue
        if t[1] == "interface-range":
            notices.append("interface-range residual; verificar expansão de herança")
            continue
        name, tail = t[1], t[2:]
        if tail[0] == "unit" and len(tail) >= 3:
            name += "." + tail[1]
            tail = tail[2:]
        row = interfaces.setdefault(name, {"interface": name})
        if tail[0] == "description":
            row["descricao"] = " ".join(tail[1:])
        if tail[0] == "disable":
            row["admin_configurado"] = "down"
        if tail[0] in ("mtu", "speed", "native-vlan-id"):
            row[{"native-vlan-id": "vlan_nativa", "speed": "velocidade_configurada"}.get(tail[0], tail[0])] = " ".join(tail[1:])
        if "802.3ad" in tail:
            group = tail[tail.index("802.3ad") + 1]
            row["agregacao"] = group
            groups.append({"agregacao": group, "membro": name})
        if "lacp" in tail:
            row["lacp"] = " ".join(tail[tail.index("lacp") + 1:])
        if tail[:2] == ["family", "ethernet-switching"]:
            rest = tail[2:]
            if rest and rest[0] in ("interface-mode", "port-mode"):
                row["modo_configurado"] = rest[1]
            if rest[:2] == ["vlan", "members"]:
                members = [x for x in rest[2:] if x not in ("[", "]")]
                row.setdefault("membros_vlan", []).extend(members)
        if "address" in tail and tail[:1] == ["family"]:
            row["enderecos"] = (row.get("enderecos", "") + "; " + tail[tail.index("address") + 1]).strip("; ")
    for row in interfaces.values():
        members = row.get("membros_vlan", [])
        resolved = [str(vlans.get(m, {}).get("vlan_id", m)) for m in members]
        if members:
            row["membros_vlan"] = ",".join(members)
            field = "vlan_access" if row.get("modo_configurado") == "access" else "vlans_membros"
            row[field] = ",".join(resolved)
    for group in groups:
        group["modo"] = interfaces.get(group["agregacao"], {}).get("lacp", "não explícito")
    return list(interfaces.values()), list(vlans.values()), groups, sorted(set(notices))


COMMANDS = {
    "cisco_ios": [("Equipamento", "show version"), ("Hardware", "show inventory"),
                  ("Estado_portas", "show interfaces status"), ("Switchport", "show interfaces switchport"),
                  ("VLANs_operacionais", "show vlan brief"), ("LAG_operacional", "show etherchannel summary"),
                  ("Vizinhos_LLDP", "show lldp neighbors detail"), ("Vizinhos_CDP", "show cdp neighbors detail"),
                  ("Enderecos_IP", "show ip interface brief")],
    "cisco_nxos": [("Equipamento", "show version"), ("Hardware", "show inventory"),
                   ("Estado_portas", "show interface status"), ("Switchport", "show interface switchport"),
                   ("VLANs_operacionais", "show vlan brief"), ("LAG_operacional", "show port-channel summary"),
                   ("Vizinhos_LLDP", "show lldp neighbors detail"), ("Vizinhos_CDP", "show cdp neighbors detail"),
                   ("Enderecos_IP", "show ip interface brief")],
    "juniper_junos": [("Equipamento", "show version"), ("Hardware", "show chassis hardware"),
                      ("Estado_portas", "show interfaces terse"), ("Detalhes_portas", "show interfaces extensive"),
                      ("VLANs_operacionais", "show vlans"), ("LAG_operacional", "show lacp interfaces"),
                      ("Vizinhos_LLDP", "show lldp neighbors")],
}

# Additional read-only queries needed by the user's port verification template.
for platform in ('cisco_ios', 'cisco_nxos'):
    COMMANDS[platform] += [('STP', 'show spanning-tree'), ('PoE', 'show power inline'),
                           ('Detalhes_portas', 'show interfaces')]
COMMANDS['juniper_junos'] += [('STP', 'show spanning-tree interface'),
                             ('LLDP_XML', 'show lldp neighbors | display xml | no-more'),
                             ('PoE', 'show poe interface'),
                             ('Detalhes_XML', 'show interfaces extensive | display xml | no-more'),
                             ('LACP_detalhes', 'show lacp interfaces extensive')]


def cli_error(output):
    return bool(re.search(r"(?im)^\s*(?:%\s*(?:Invalid|Error|Ambiguous|Incomplete|Unknown|Authorization)|error:|syntax error|permission denied|unknown command)|ACCESS-DENIED", output))


def ssh_diagnostic(exc, credentials):
    message = str(exc)
    for key in ('password', 'secret'):
        value = credentials.get(key)
        if value:
            message = message.replace(value, '[oculto]')
    message = re.sub(r'[\x00-\x1f\x7f]', ' ', message).strip()[:1500]
    lower = message.lower()
    if 'known_hosts' in lower or 'not found in known' in lower:
        hint = 'Chave SSH desconhecida: conecte com OpenSSH e confira a impressão digital antes de cadastrar.'
    elif 'host key' in lower or 'hostkey' in lower:
        hint = 'Falha na verificação da chave SSH; confira a impressão digital com o administrador.'
    elif 'authentication' in lower or 'authentication' in type(exc).__name__.lower():
        hint = 'Autenticação recusada: confira usuário, senha e permissão de acesso SSH.'
    else:
        hint = 'Confira IP, porta SSH, conectividade/VPN e a mensagem abaixo.'
    return f'{type(exc).__name__}: {hint} Detalhe: {message or "sem detalhe adicional"}'


def collect(device, credentials, timeout, known_hosts):
    from refresh_core.ssh import connect_switch as ConnectHandler
    from netmiko.utilities import get_structured_data_textfsm
    tables, errors = {}, []
    stamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    meta = {"equipamento": device["nome"], "host": device["host"], "coleta_utc": stamp}
    def add(sheet, rows, source):
        tables.setdefault(sheet, []).extend([{**meta, **r, "fonte": source} for r in rows])
    def issue(command, message):
        errors.append({**meta, "comando": command, "erro": message})
    parameters = dict(device_type=device["plataforma"], host=device["host"], port=int(device["porta"] or 22),
                      username=credentials["username"], password=credentials["password"],
                      secret=credentials["secret"], conn_timeout=timeout, auth_timeout=timeout,
                      banner_timeout=timeout, ssh_strict=True, system_host_keys=True)
    if known_hosts:
        parameters.update(alt_host_keys=True, alt_key_file=str(known_hosts))
    if device.get("key_file"):
        parameters.update(use_keys=True, key_file=device["key_file"])
    config_ok = False
    try:
        with ConnectHandler(**parameters) as conn:
            if device["plataforma"].startswith("cisco") and credentials["secret"]:
                conn.enable()
            command = ("show configuration | display inheritance | display set | no-more"
                       if device["plataforma"] == "juniper_junos" else "show running-config")
            try:
                config = conn.send_command(command, read_timeout=timeout)
                if cli_error(config) or not config.strip():
                    raise ValueError("Configuração indisponível ou acesso negado")
                if device["plataforma"] == "juniper_junos":
                    rows, vlans, groups, notes = parse_junos(config)
                else:
                    rows, vlans, groups, notes = parse_cisco(config)
                add("Interfaces_config", rows, command)
                add("VLANs_config", vlans, command)
                add("Agregacoes_config", groups, command)
                config_ok = True
                if not rows:
                    issue(command, "Nenhuma interface extraída; verificar formato/permissões da configuração")
                for note in notes:
                    issue(command, note)
            except Exception as exc:
                issue(command, type(exc).__name__ + ": configuração não coletada/interpretada")
            commands = list(COMMANDS[device['plataforma']])
            lldp_ports = set()
            for sheet, command in commands:
                try:
                    raw = conn.send_command(command, read_timeout=timeout)
                    add("Evidencias", [{"comando": command, "parte": i // 30000 + 1, "saida": raw[i:i+30000]}
                                       for i in range(0, len(raw), 30000)], command)
                    if cli_error(raw):
                        raise ValueError("Comando indisponível ou acesso negado")
                    from operacional import parse
                    parsed = parse(device['plataforma'], 'Vizinhos_LLDP' if sheet == 'LLDP_porta' else sheet, raw)
                    if not parsed:
                        parsed = get_structured_data_textfsm(raw, platform=device["plataforma"], command=command)
                    if isinstance(parsed, list) and parsed:
                        add({'Detalhes_XML': 'Detalhes_portas', 'LLDP_XML': 'Vizinhos_LLDP', 'LLDP_porta': 'Vizinhos_LLDP'}.get(sheet, sheet), parsed, command)
                        if device['plataforma'] == 'juniper_junos' and sheet in ('Vizinhos_LLDP', 'LLDP_XML'):
                            for neighbor in parsed:
                                port = neighbor.get('local_interface') or neighbor.get('local_port', '')
                                if re.fullmatch(r'[a-zA-Z][a-zA-Z0-9/.:_-]*', str(port)) and port not in lldp_ports:
                                    lldp_ports.add(port)
                                    commands.append(('LLDP_porta', f'show lldp neighbors interface {port} | no-more'))
                    else:
                        issue(command, "Sem dados estruturados: saída vazia, tabela vazia ou template indisponível. Consulte Evidencias.")
                except Exception as exc:
                    message = ssh_diagnostic(exc, credentials).split(' Detalhe: ', 1)[-1]
                    issue(command, type(exc).__name__ + ': falha na coleta/interpretação: ' + message)
    except Exception as exc:
        issue("SSH", ssh_diagnostic(exc, credentials))
    has_data = config_ok or any(v for k, v in tables.items() if k != 'Evidencias')
    add("Resumo", [{"plataforma": device["plataforma"], "status": "PARCIAL" if errors and has_data else "FALHA" if errors else "OK",
                     "config_coletada": config_ok, "interfaces_config": len(tables.get("Interfaces_config", [])),
                     "vlans_config": len(tables.get("VLANs_config", [])), "ocorrencias": len(errors)}], "coletor")
    tables["Ocorrencias"] = errors
    return tables


def export_excel(tables, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    order = ["Resumo", "Interfaces_config", "VLANs_config", "Agregacoes_config"]
    names = order + sorted(set(tables) - set(order) - {"Evidencias", "Ocorrencias"}) + ["Ocorrencias", "Evidencias"]
    for index, name in enumerate(names):
        rows = tables.get(name, [])
        ws = wb.create_sheet(name[:31])
        columns = list(dict.fromkeys(key for row in rows for key in row)) or ["informacao"]
        ws.append(columns)
        for row in rows:
            values = []
            for column in columns:
                value = row.get(column, "")
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                if isinstance(value, str):
                    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)[:32767]
                values.append(value)
            ws.append(values)
            for cell in ws[ws.max_row]:
                # All device-provided strings are text, including leading =, +, -, @.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="17365D")
        ws.freeze_panes = "D2"
        ws.sheet_view.showGridLines = False
        if rows:
            table = Table(displayName=f"Dados{index}", ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        for i, column in enumerate(columns, 1):
            length = max([len(column)] + [len(str(r.get(column, ""))) for r in rows[:100]])
            ws.column_dimensions[get_column_letter(i)].width = min(65, max(16, length + 2))
        if name == "Evidencias":
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                ws.row_dimensions[row[0].row].height = 60
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.xlsx")
    wb.save(temporary)
    temporary.replace(path)


def check_dependencies():
    for module in ("netmiko", "ntc_templates", "openpyxl"):
        try:
            importlib.import_module(module)
        except ImportError as exc:
            requirements = Path(__file__).resolve().with_name("requirements.txt")
            print(f"Dependência indisponível ao carregar {module}: {exc}", file=sys.stderr)
            print("Instale as dependências no mesmo Python usado para executar o programa.", file=sys.stderr)
            print(f'PowerShell: & "{sys.executable}" -m pip install -r "{requirements}"', file=sys.stderr)
            return False
    return True


def read_inventory(path):
    with open(path, encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        if not {'nome', 'host', 'plataforma'} <= set(reader.fieldnames or []):
            raise ValueError('Inventário precisa das colunas nome,host,plataforma, separadas por vírgulas')
        devices = []
        names, endpoints = set(), set()
        for number, row in enumerate(reader, 2):
            if None in row:
                raise ValueError(f'Linha {number}: colunas extras; confira as vírgulas')
            d = {k: (v or '').strip() for k, v in row.items()}
            if not any(d.values()):
                continue
            if not d['nome'] or not d['host'] or d['plataforma'] not in PLATFORMS:
                raise ValueError(f'Linha {number}: nome/host vazio ou plataforma inválida')
            port = d.get('porta') or '22'
            if not port.isdigit() or not 1 <= int(port) <= 65535:
                raise ValueError(f'Linha {number}: porta inválida')
            d['porta'] = port
            endpoint = (d['host'].casefold(), int(port))
            if d['nome'].casefold() in names or endpoint in endpoints:
                raise ValueError(f'Linha {number}: nome ou endereço/porta duplicado')
            names.add(d['nome'].casefold())
            endpoints.add(endpoint)
            if d.get('key_file'):
                key = Path(d['key_file']).expanduser()
                if not key.is_absolute():
                    key = Path(path).resolve().parent / key
                if not key.is_file():
                    raise ValueError(f'Linha {number}: arquivo de chave não encontrado')
                d['key_file'] = str(key)
            devices.append(d)
    if not devices:
        raise ValueError('Inventário vazio')
    return devices


def run_mapping(jobs, args, on_event=None):
    completed = 0
    def notify(message):
        if on_event: on_event({'message': message, 'completed': completed, 'total': len(jobs)})
        else: print(message)
    notify('Consultando switches por SSH…')
    result = {}
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = {pool.submit(collect, d, c, args.timeout, args.known_hosts): d for d, c in jobs}
        for future in as_completed(futures):
            try:
                data = future.result()
            except Exception as exc:
                d = futures[future]
                meta = {'equipamento': d['nome'], 'host': d['host'],
                        'coleta_utc': datetime.now(timezone.utc).isoformat(timespec='seconds')}
                data = {'Resumo': [{**meta, 'status': 'FALHA'}],
                        'Ocorrencias': [{**meta, 'comando': 'Coletor',
                                         'erro': type(exc).__name__ + ': falha inesperada; demais switches continuam'}]}
            for name, rows in data.items():
                result.setdefault(name, []).extend(rows)
            completed += 1
            notify(f"{futures[future]['nome']}: {data['Resumo'][0]['status']}")
            for occurrence in data.get('Ocorrencias', []):
                if occurrence.get('comando') in ('SSH', 'Coletor'):
                    notify('  ' + occurrence['erro'])
    for rows in result.values():
        rows.sort(key=lambda r: (r.get("equipamento", ""), r.get("interface", "")))
    target = Path(args.saida) if args.saida else Path(__file__).parent / 'relatorios' / f"mapeamento_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
    if args.modelo:
        from modelo_excel import export_model
        export_model(result, args.modelo, target)
    else:
        export_excel(result, target)
    notify(f"Excel: {target.resolve()}")
    if args.diagnostico:
        target.with_suffix('.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    from refresh_core.storage import event
    event('Mapeamento de portas',target,status='Parcial' if result.get('Ocorrencias') else 'Concluído',detail=str(len(jobs))+' switches')
    return {"path": target.resolve(), "tables": result, "code": 2 if result.get("Ocorrencias") else 0}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--inventario", default=str(Path(__file__).with_name('inventario.csv')))
    parser.add_argument("--saida", default=None)
    parser.add_argument("--modelo", type=Path, default=Path(__file__).with_name('modelo_fixo.xlsx'), help="Modelo fixo; original preservado")
    parser.add_argument("--diagnostico", action='store_true', help="Salvar evidências operacionais em JSON (opcional)")
    parser.add_argument("--usuario", default=os.getenv("SW_USER", ""))
    parser.add_argument("--known-hosts", type=Path)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--timeout", type=int, default=90)
    parser.add_argument("--nao-interativo", action="store_true")
    args = parser.parse_args()
    if not check_dependencies():
        return 1
    if args.modelo:
        from modelo_excel import load_model
        try:
            load_model(args.modelo).close()
        except (OSError, ValueError) as exc:
            parser.error(str(exc))
        if args.saida and Path(args.saida).resolve() == args.modelo.resolve():
            parser.error("--saida não pode sobrescrever --modelo")
    if not 1 <= args.workers <= 32 or args.timeout < 1:
        parser.error("workers deve ser 1..32 e timeout positivo")
    try:
        devices = read_inventory(args.inventario)
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
    if args.saida and Path(args.saida).suffix.lower() != '.xlsx':
        parser.error('--saida precisa terminar em .xlsx')
    secrets = {}
    jobs = []
    for d in devices:
        if d["plataforma"] not in PLATFORMS or not d["host"] or not d["nome"]:
            parser.error("Equipamento incompleto ou plataforma inválida")
        d.setdefault("porta", "22")
        if not (d["porta"] or "22").isdigit() or not 1 <= int(d["porta"] or 22) <= 65535:
            parser.error("Porta inválida")
        user = d.get("usuario") or args.usuario
        if not user and not args.nao_interativo:
            args.usuario = user = input("Usuário SSH: ").strip()
        if not user:
            parser.error("Informe --usuario, SW_USER ou usuario no CSV")
        def credential(field, fallback, optional=False):
            env = d.get(field) or fallback
            if not env:
                return ""
            if env in os.environ:
                return os.environ[env]
            if env not in secrets:
                if args.nao_interativo:
                    if optional:
                        return ""
                    parser.error(f"Variável {env} não definida")
                secrets[env] = getpass.getpass(f"Credencial {env} (não será salva): ")
            return secrets[env]
        password = "" if d.get("key_file") else credential("password_env", "SW_PASSWORD")
        secret = credential("secret_env", "", optional=False)
        jobs.append((d, {"username": user, "password": password, "secret": secret}))
    return run_mapping(jobs, args)['code']


if __name__ == "__main__":
    if len(sys.argv) == 1:
        from interface import main as gui_main
        gui_main()
    else:
        sys.exit(main())
