import copy
import json
import queue
import re
import threading
from pathlib import Path
import tkinter as tk
from tkinter import ttk,filedialog,messagebox,simpledialog
import hardening as engine
import evolucao
import reversao

class InteractiveAuth:
    def __init__(self,password,ask):self.password=password;self.ask=ask
    def respond(self,title,instructions,prompts):
        replies=[]
        for prompt,echo in prompts:
            if re.search(r'(?i)token|otp|verification|code',prompt):
                reply=self.ask(title)
                if not reply or not re.fullmatch(r'[0-9]{6,10}',reply):raise ValueError('OTP cancelado ou inválido')
                replies.append(reply)
            elif re.search(r'(?i)password|senha',prompt):replies.append(self.password)
            else:raise ValueError('Desafio SSH não reconhecido')
        return replies

class Extras:
    def install_extras(self,policy,policy_tabs,connection):
        self.otp_enabled=tk.BooleanVar(value=False)
        ttk.Checkbutton(connection,text='SSH com OTP (solicitar código a cada nova conexão)',variable=self.otp_enabled).grid(row=5,column=0,columnspan=3,sticky='w')
        self.backup_secret=tk.StringVar()
        self.exceptions=copy.deepcopy(policy.get('exceptions',[]))
        page=ttk.Frame(policy_tabs,padding=10);policy_tabs.add(page,text='Exceções')
        ttk.Label(page,text='Exceções mantêm a divergência visível e excluem a correção. Serial obrigatório; objeto vazio abrange toda a regra.').pack(anchor='w')
        self.exception_tree=ttk.Treeview(page,columns=('serial','rule','entry','reason','expires'),show='headings',height=8)
        for key,title in [('serial','Serial'),('rule','Regra'),('entry','Conta / interface'),('reason','Justificativa'),('expires','Validade AAAA-MM-DD')]:self.exception_tree.heading(key,text=title);self.exception_tree.column(key,width=150)
        self.exception_tree.pack(fill='both',expand=True)
        row=ttk.Frame(page);row.pack(fill='x');self.exception_vars={}
        for i,key in enumerate(('serial','rule','entry','reason','expires')):
            v=tk.StringVar();self.exception_vars[key]=v
            widget=ttk.Combobox(row,textvariable=v,values=list(engine.LABELS),state='readonly') if key=='rule' else ttk.Entry(row,textvariable=v)
            widget.grid(row=0,column=i,sticky='ew');row.columnconfigure(i,weight=1)
        ttk.Button(page,text='Adicionar exceção',command=self.add_exception).pack(side='left',padx=5)
        ttk.Button(page,text='Remover selecionadas',command=self.remove_exception).pack(side='left')
        self.draw_exceptions()
        page=ttk.Frame(policy_tabs,padding=10);policy_tabs.add(page,text='Compatibilidade')

        self.token_text=tk.Text(page,height=5);self.token_text.insert('1.0','\n'.join(k+'='+v for k,v in policy['correcoes'].get('mfa_tokens',{}).items()))
        ttk.Label(page,text='Perfis: 40F, 60F e 100F / 7.4.9 e 7.4.12; 50E / 6.2.16. Campos observados confirmam disponibilidade; ausências não são presumidas como padrão.\nMaintainer: aplicável somente ao 50E 6.2.16. Senha desativada: comprimento pode ficar oculto. Outros modelos: suporte condicionado à evidência coletada.',wraplength=900).pack(anchor='w',pady=12)
        ttk.Button(page,text='Consultar compatibilidade da última coleta',command=self.show_compatibility).pack(anchor='w')
        page=ttk.Frame(policy_tabs,padding=12);policy_tabs.add(page,text='Histórico e recuperação')
        ttk.Label(page,text='Cada Excel acompanha um JSON de verificação. Comparações usam serial e regra; mudanças de política são identificadas.',wraplength=900).pack(anchor='w',pady=8)
        for label,callback in [('Comparar duas verificações',self.compare_reports),('Preparar reversão de uma aplicação',self.prepare_rollback),('Exportar backup descriptografado',self.export_backup)]:ttk.Button(page,text=label,command=callback).pack(anchor='w',pady=8)
        ttk.Label(page,text='Backup SCP exige admin-scp habilitado e permissão de leitura completa. A automação não habilita esse serviço silenciosamente.\nA restauração completa é manual e pode reiniciar o equipamento. Senha do backup não é salva.',wraplength=900).pack(anchor='w',pady=12)

    def ensure_backup_password(self,parent):
        if len(self.backup_secret.get())>=12:return True
        value=simpledialog.askstring('Proteger backup','Defina uma senha de 12 ou mais caracteres para proteger os backups.\nGuarde-a para restaurar. Será solicitada apenas nesta primeira aplicação da sessão.',show='*',parent=parent)
        if value is None:return False
        if len(value)<12:messagebox.showerror('Backup','Use pelo menos 12 caracteres.',parent=parent);return False
        self.backup_secret.set(value);return True

    def draw_exceptions(self):
        self.exception_tree.delete(*self.exception_tree.get_children())
        for i,e in enumerate(self.exceptions):self.exception_tree.insert('','end',iid=str(i),values=[e[k] for k in ('serial','rule','entry','reason','expires')])
    def add_exception(self):
        if self.busy:return
        try:
            rows=self.exceptions+[{k:v.get().strip() for k,v in self.exception_vars.items()}]
            self.exceptions=evolucao.validate_exceptions(rows,engine.LABELS);self.draw_exceptions()
        except Exception as exc:messagebox.showerror('Exceção',str(exc))
    def remove_exception(self):
        if self.busy:return
        selected={int(i) for i in self.exception_tree.selection()};self.exceptions=[e for i,e in enumerate(self.exceptions) if i not in selected];self.draw_exceptions()
    def token_map(self):
        values={}
        for line in self.token_text.get('1.0','end-1c').splitlines():
            if not line.strip():continue
            if '=' not in line:raise ValueError('FortiToken: use conta=serial')
            k,v=(x.strip() for x in line.split('=',1))
            if k in values:raise ValueError('Conta FortiToken repetida')
            values[k]=v
        return values
    def ask_otp(self,title=''):
        response=queue.Queue();self.events.put(('otp',(response,title)))
        try:return response.get(timeout=120)
        except queue.Empty:return None
    def show_text(self,title,text):
        win=tk.Toplevel(self.root);win.title(title);win.geometry('950x600')
        box=tk.Text(win,wrap='word');box.pack(fill='both',expand=True);box.insert('1.0',text);box.configure(state='disabled')
        return win
    def show_compatibility(self):
        if not self.report:return messagebox.showinfo('Compatibilidade','Execute uma verificação primeiro')
        self.show_text('Compatibilidade','\n'.join(d['device']['nome']+' / '+evolucao.model(d['snapshot'])+' / '+d['identity']['version']+' / '+f['title']+': '+f.get('compatibility','Não informado') for d in self.report['devices'] for f in d['findings']))
    def compare_reports(self):
        if self.busy:return
        try:
            a=filedialog.askopenfilename(title='Verificação anterior',filetypes=[('JSON','*.json')])
            if not a:return
            b=filedialog.askopenfilename(title='Verificação atual',filetypes=[('JSON','*.json')])
            if not b:return
            rows=evolucao.compare(json.loads(Path(a).read_text(encoding='utf-8')),json.loads(Path(b).read_text(encoding='utf-8')))
            from openpyxl import Workbook
            path=filedialog.asksaveasfilename(title='Salvar comparação',defaultextension='.xlsx',initialfile='comparativo_hardening.xlsx')
            if not path:return
            wb=Workbook();ws=wb.active;ws.title='Comparação';ws.append(['Equipamento','Regra','Antes','Depois','Resultado','Política mudou'])
            for row in rows:ws.append(list(row.values()))
            for row in ws:
                for c in row:
                    if isinstance(c.value,str):c.data_type='s'
            from openpyxl.utils import get_column_letter
            for i in range(1,7):ws.column_dimensions[get_column_letter(i)].width=32
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;wb.save(path)
            self.show_text('Comparação salva','\n'.join(str(r) for r in rows))
        except Exception as exc:messagebox.showerror('Comparação',str(exc))
    def export_backup(self):
        if self.busy:return
        try:
            path=filedialog.askopenfilename(filetypes=[('Backup criptografado','*.fgbackup')])
            if not path:return
            password=simpledialog.askstring('Backup','Senha de criptografia do backup:',show='*',parent=self.root)
            if not password:return
            data=engine.backup_seguro.decrypt(path,password)
            target=filedialog.asksaveasfilename(title='Exportar configuração com segredos — guarde em local protegido',defaultextension='.conf')
            if target:Path(target).write_bytes(data);self.status.set('Backup exportado; confira o procedimento .restauracao.txt')
        except Exception:messagebox.showerror('Backup','Não foi possível descriptografar/exportar. Confira senha, arquivo e destino.')
    def prepare_rollback(self):
        if self.busy:return
        try:
            if not self.report:raise ValueError('Execute nova verificação dos equipamentos antes de preparar a reversão')
            path=filedialog.askopenfilename(title='Diário da aplicação original',filetypes=[('Diário','*.jsonl')])
            if not path:return
            plan=reversao.prepare(path,self.report);timeout,_,known,password=self.settings()
            if not self.ensure_backup_password(self.root):return
            backup_password=self.backup_secret.get()
            win=self.show_text('Revisão da reversão','Reverter pode restaurar configurações menos seguras. Confira cada campo. Novo backup será feito.\n\n'+'\n\n'.join(i['device']['nome']+'\n'+'\n'.join(o['field']+': '+o['before']+' → '+o['after'] for o in i['operations'])+'\n'+'\n'.join(line for batch in reversao.batches(i['operations']) for line in batch) for i in plan))
            win.transient(self.root);win.grab_set()
            confirmation=tk.StringVar();ttk.Label(win,text='Digite REVERTER para aprovar este plano:').pack();ttk.Entry(win,textvariable=confirmation).pack()
            def confirm():
                if self.busy:return
                if confirmation.get()!='REVERTER':return
                win.destroy();self.set_busy(True);self.cancel.clear();self.cancel_button.configure(state='normal');token=engine.digest(plan)
                def work():
                    try:
                        rows,journal=reversao.apply(plan,token,password,backup_password,timeout,known,engine.BASE/'resultados',self.cancel)
                        self.events.put(('rollback',(rows,journal)))
                    except Exception as exc:self.events.put(('error','Reversão: '+type(exc).__name__))
                threading.Thread(target=work,daemon=True).start()
            ttk.Button(win,text='Aprovar reversão',command=confirm).pack()
        except Exception as exc:messagebox.showerror('Reversão',str(exc))
