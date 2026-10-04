"""Auditoria SSH e remediações restritas. Nenhuma escrita durante a auditoria."""
import copy
import time
import evolucao
import backup_seguro
import remediacao
import hashlib
import ipaddress
import json
import re
import shlex
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from fortigate_ssh import SSH, COMMANDS
from fortios_parser import parse, config_blocks, commands as cli_commands
from padrao_empresa import RULES, TITLES, PRE, POST, MAP, EXTRA_AUTO, normalize_banner

BASE=Path(__file__).resolve().parent
UNKNOWN='Não foi possível verificar'
READS={'Sistema':'get system status','HA':'get system ha status',
 'Global':'show full-configuration system global',
 'Interfaces':'show full-configuration system interface',
 'Administradores':'show full-configuration system admin',
 'Senha':'show full-configuration system password-policy',
 'NTP':'show full-configuration system ntp',
 'SNMP':'show full-configuration system snmp community'}
READS.update({'FortiAnalyzer':'show full-configuration log fortianalyzer setting','AutoInstall':'show full-configuration system auto-install','Eventos':'show full-configuration log eventfilter','BannerPre':'show full-configuration system replacemsg admin pre_admin-disclaimer-text','BannerPost':'show full-configuration system replacemsg admin post_admin-disclaimer-text'})
READS['Session']='get system admin status'
READS['Tokens']='show full-configuration user fortitoken'
READS['TokenUsers']='show user local'
READS['Email']='show full-configuration system email-server'
for removed in ('Administradores','Senha','SNMP','Tokens','TokenUsers','Email'):READS.pop(removed,None)
COMMANDS.update(READS)
GETS={cat:cmd.replace('show full-configuration','get') for cat,cmd in READS.items() if cat in ('Global','Senha','FortiAnalyzer','NTP','AutoInstall','Eventos')}
COMMANDS.update({'get_'+cat:cmd for cat,cmd in GETS.items()})
ERROR=re.compile(r'(?im)(command fail|command parse error|permission denied|unknown action|return code\s*-|object check operator error|entry not found|value parse error|incomplete command|ambiguous command)')
GLOBAL_FIELDS={'admintimeout','admin-lockout-duration','admin-https-ssl-versions','admin-https-redirect','cfg-save'}
SIMPLE_FIELDS={'Email':{'server'},'Global':GLOBAL_FIELDS,'Senha':{'status','minimum-length'},'NTP':{'ntpsync','type'}}
BLOCK_FIELDS={'Interfaces':{'allowaccess','ip','role','status'},'Administradores':{'accprofile','remote-auth','two-factor','email-to','fortitoken'}|{f'trusthost{i}' for i in range(1,11)},'SNMP':{'status','query-v1-status','query-v2c-status','trap-v1-status','trap-v2c-status'}}
LABELS={'timeout':'Timeout administrativo','tentativas':'Limite de tentativas administrativas','bloqueio':'Duração do bloqueio administrativo',
 'cripto':'Criptografia forte','tls':'Versões TLS da administração','http_telnet':'HTTP/Telnet nas interfaces',
 'wan_admin':'Acesso administrativo em interface WAN','trusted_hosts':'Restrição IPv4 dos administradores',
 'mfa':'Segundo fator nos administradores locais','senha':'Política de senha local','ntp':'Sincronização NTP configurada','snmp':'SNMP v1/v2c'}
LABELS.update(TITLES)
SCRIPT_IDS=set(RULES)|{'timeout','bloqueio','tls','ntp'}
LEGACY_LABELS=LABELS.copy()
LABELS={k:v for k,v in LABELS.items() if k in SCRIPT_IDS}
DEFAULT_POLICY={'enabled':list(LABELS),'timeout_max':5,'tentativas_max':3,'bloqueio_min':600,'senha_min':12,'mapa_interfaces':MAP.copy(),'banners':{'pre':PRE,'post':POST},'tls_permitidos':['tlsv1-2']}
DEFAULT_POLICY['exceptions']=[]
DEFAULT_POLICY['correcoes']=copy.deepcopy(remediacao.DEFAULTS)
for section,field,target,reason in RULES.values():
    if section=='Interfaces':BLOCK_FIELDS['Interfaces'].add(field)
    else:SIMPLE_FIELDS.setdefault(section,set()).add(field)
SIMPLE_FIELDS['FortiAnalyzer'].add('status')
AUTO={'timeout':('admintimeout','timeout_max',1,30),'tentativas':('admin-lockout-threshold','tentativas_max',1,10),'bloqueio':('admin-lockout-duration','bloqueio_min',60,900)}
IMPACT={'timeout':'Sessões administrativas ociosas poderão expirar mais cedo.',
 'tentativas':'Erros repetidos de autenticação poderão bloquear administradores; confira as credenciais das integrações.',
 'bloqueio':'Administradores que excederem as tentativas precisarão aguardar o período de bloqueio.'}
SUPPORTED={'7.4.9','7.4.12'}

def application_reason(f):
    if f['eligible']:return 'Marque este item para revisar os comandos e aprovar a aplicação.'
    if f['status']=='Conforme':return 'O controle já atende à política; não há alteração a aplicar.'
    if f['status']=='Não aplicável':return f['evidence']
    if f['status']==UNKNOWN:return 'A coleta não comprovou o estado. Corrija a consulta e audite novamente antes de alterar.'
    if f['blocked']:return f['blocked']
    return 'Nenhuma alteração concreta gerada. Confira a evidência e os parâmetros da política antes de verificar novamente.'

def application_label(f):
    if f['eligible']:return 'Pode aplicar após aprovação'
    if f['status']=='Conforme':return 'Nenhuma alteração necessária'
    if f['status']=='Não aplicável':return 'Não aplicável nesta versão'
    if f['status']==UNKNOWN:return 'Resolver coleta primeiro'
    return 'Bloqueado: ver detalhe' if f['blocked'] else 'Revisar parâmetros'

def validate_policy(p):
    p=copy.deepcopy(p)
    if isinstance(p,dict):
        p.setdefault('exceptions',[])
        if isinstance(p.get('enabled'),list):
            if any(k not in LEGACY_LABELS for k in p['enabled']):raise ValueError('Regra desconhecida')
            p['enabled']=[k for k in p['enabled'] if k in SCRIPT_IDS] or list(LABELS)
        if isinstance(p.get('exceptions'),list):p['exceptions']=[e for e in p['exceptions'] if e.get('rule') in SCRIPT_IDS]
    if not isinstance(p,dict) or set(p)!=set(DEFAULT_POLICY): raise ValueError('Política de hardening inválida')
    if not isinstance(p['enabled'],list) or not p['enabled'] or any(x not in LABELS for x in p['enabled']) or len(set(p['enabled']))!=len(p['enabled']): raise ValueError('Selecione regras válidas, sem repetições')
    for key,lo,hi in [('timeout_max',1,30),('tentativas_max',1,10),('bloqueio_min',60,900),('senha_min',8,128)]:
        if type(p[key]) is not int or not lo<=p[key]<=hi: raise ValueError(f'{key}: informe inteiro entre {lo} e {hi}')
    if not isinstance(p['mapa_interfaces'],dict) or set(p['mapa_interfaces'])!=set(MAP):raise ValueError('Mapeie as cinco interfaces do script')
    if any(not isinstance(v,str) or not re.fullmatch(r'[A-Za-z0-9_.:/-]{1,79}',v) for v in p['mapa_interfaces'].values()):raise ValueError('Nome de interface inválido')
    if len(set(p['mapa_interfaces'].values()))!=5:raise ValueError('Cada porta do script deve ter um destino distinto')
    if not isinstance(p['banners'],dict) or set(p['banners'])!={'pre','post'}:raise ValueError('Banners inválidos')
    if any(not isinstance(v,str) or not v.strip() or len(v)>16000 or '\x00' in v for v in p['banners'].values()):raise ValueError('Texto dos banners inválido')
    if not isinstance(p['tls_permitidos'],list) or not p['tls_permitidos'] or any(x not in ('tlsv1-2','tlsv1-3') for x in p['tls_permitidos']):raise ValueError('TLS permitido: tlsv1-2 e/ou tlsv1-3')
    evolucao.validate_exceptions(p['exceptions'],LABELS)
    p['correcoes']=remediacao.validate(p['correcoes'])
    return copy.deepcopy(p)



def digest(data): return hashlib.sha256(json.dumps(data,sort_keys=True,ensure_ascii=False).encode()).hexdigest()

def migrate_policy(p):
    p=copy.deepcopy(p)
    if isinstance(p,dict) and 'mapa_interfaces' in p:p.setdefault('exceptions',[])
    old={'enabled','timeout_max','tentativas_max','bloqueio_min','senha_min'}
    if isinstance(p,dict) and set(p)==old:
        new=copy.deepcopy(DEFAULT_POLICY)
        new.update(p)
        new['enabled']=list(dict.fromkeys(p['enabled']+list(RULES)))
        # Valores padrão anteriores são atualizados para o script recebido.
        if p['timeout_max']==10:new['timeout_max']=5
        if p['bloqueio_min']==60:new['bloqueio_min']=600
        return validate_policy(new)
    if isinstance(p,dict) and set(p)==set(DEFAULT_POLICY)-{'correcoes'}:
        p=copy.deepcopy(p);p['correcoes']=copy.deepcopy(remediacao.DEFAULTS)
    return validate_policy(p)

def read_config(category,text):
    header=READS[category].replace('show full-configuration','config')
    if shlex.split(text.strip().splitlines()[0])!=shlex.split(header) or not text.rstrip().endswith('end'): raise ValueError('Configuração incompleta')
    stack=[]
    for raw in cli_commands(text):
        line=raw.strip()
        if line.startswith('config '):stack.append('config')
        elif line.startswith('edit '):
            if not stack or stack[-1]!='config':raise ValueError('Estrutura de configuração inválida')
            stack.append('edit')
        elif line in ('end','next'):
            expected='config' if line=='end' else 'edit'
            if not stack or stack.pop()!=expected:raise ValueError('Configuração incompleta')
    if stack:raise ValueError('Configuração incompleta')
    # Somente campos necessários sobrevivem à coleta; hashes e comunidades não são armazenados.
    if category in SIMPLE_FIELDS:
        values={}; depth=0
        for raw in cli_commands(text):
            line=raw.strip()
            if line.startswith('config '): depth+=1
            elif line=='end': depth-=1
            elif depth==1 and line.startswith(('set ','unset ')):
                parts=shlex.split(line)
                if len(parts)>=2 and parts[1] in SIMPLE_FIELDS[category]:
                    if parts[1] in values: raise ValueError('Campo repetido')
                    values[parts[1]]=' '.join(parts[2:]) if parts[0]=='set' else ''
        if depth!=0: raise ValueError('Configuração incompleta')
        return values
    rows=[]
    for n,block in enumerate(config_blocks(text),1):
        values={k:v for k,v in block['settings'].items() if k in BLOCK_FIELDS[category]}
        values['id']=block['id']
        rows.append(values)
    return rows

def snapshot(connection):
    result={}
    for category,command in READS.items():
        try:text=connection.command(command)
        except (TimeoutError,ConnectionError,OSError) as exc:
            result[category]={'error':'Falha durante consulta: '+type(exc).__name__}
            for remaining in READS:
                if remaining not in result:result[remaining]={'error':'Coleta interrompida após falha na sessão SSH'}
            return result
        if ERROR.search(text):
            result[category]={'error':evolucao.error_kind(text),'kind':evolucao.error_kind(text)}
        else:
            try:
                if category=='Session':
                    result[category]={'value':evolucao.session(text)}
                elif category=='TokenUsers':
                    if not text.strip().startswith('config user local') or not text.rstrip().endswith('end'):raise ValueError('Usuários incompletos')
                    result[category]={'value':[{'id':b['id'],'fortitoken':b['settings'].get('fortitoken','')} for b in config_blocks(text)]}
                elif category=='Tokens':
                    if not text.strip().startswith('config user fortitoken') or not text.rstrip().endswith('end'):raise ValueError('Lista de tokens incompleta')
                    result[category]={'value':[{'id':b['id'],'status':b['settings'].get('status','unknown')} for b in config_blocks(text)]}
                elif category in ('Sistema','HA'):
                    parsed=parse(category,text)
                    if parsed['status']!='ok': raise ValueError('Formato não reconhecido')
                    result[category]={'value':{r['id']:r['value'] for r in parsed['rows'] if r['id']!='System time'}}
                else:
                    values=read_config(category,text)
                    if category in GETS and set(values)!=SIMPLE_FIELDS[category]:
                        extra=connection.command(GETS[category])
                        if not ERROR.search(extra):
                            for line in extra.splitlines():
                                match=re.fullmatch(r'\s*([a-z0-9-]+)\s*:\s*(.*?)\s*',line)
                                if match and match[1] in SIMPLE_FIELDS[category] and match[1] not in values:values[match[1]]=match[2]
                    result[category]={'value':values}
            except (ValueError,IndexError,KeyError): result[category]={'error':'Formato incompleto ou não reconhecido'}
        if category=='Sistema':
            if result[category].get('value',{}).get('Virtual domain configuration','').lower()!='disable':
                return {'Sistema':{'error':'Escopo sem VDOM não confirmado'}}
    return result

def identity(s):
    system=s.get('Sistema',{}).get('value',{})
    match=re.search(r'\bv(\d+\.\d+\.\d+)\b',system.get('Version',''))
    return {'serial':system.get('Serial-Number',''),'version':match[1] if match else '',
            'hostname':system.get('Hostname',''),'mode':system.get('Current HA mode','').lower().split(',')[0].strip()}

def safety(s):
    ident=identity(s)
    if not ident['serial'] or not ident['hostname']: return 'Identidade incompleta'
    if not evolucao.legacy_50e(s) and (ident['version'] not in SUPPORTED or evolucao.model(s)=='50E'): return 'Aplicação habilitada para FortiOS 7.4.9/7.4.12 ou FortiGate 50E com 6.2.16'
    if s.get('Sistema',{}).get('value',{}).get('Virtual domain configuration','').lower()!='disable': return 'VDOM não confirmado como desabilitado'
    if s.get('Global',{}).get('value',{}).get('cfg-save')!='automatic': return 'Modo de salvamento automático não confirmado'
    if ident['mode']=='standalone': return ''
    if ident['mode'] not in ('ha a-p','ha a-a','a-p','a-a'): return 'Modo HA não reconhecido'
    ha=s.get('HA',{}).get('value',{})
    if ha.get('HA Health Status','').lower()!='ok': return 'Saúde HA não confirmada'
    members={k[7:]:v for k,v in ha.items() if k.startswith('member/')}
    if len(members)<2: return 'Menos de dois membros HA identificados'
    if members.get(ident['serial']) not in ('primary','master'): return 'Conecte ao membro primário do cluster'
    if sum(v in ('primary','master') for v in members.values())!=1: return 'Primário HA ambíguo'
    if any(ha.get('sync/'+serial)!='in-sync' for serial in members): return 'Sincronismo de todos os membros não confirmado'
    return ''

def group_id(s):
    ha=s.get('HA',{}).get('value',{})
    return tuple(sorted(k[7:] for k in ha if k.startswith('member/'))) or (identity(s)['serial'],)

def suggest_interfaces(s,current):
    names={r['id'] for r in s.get('Interfaces',{}).get('value',[])}
    result=current.copy()
    aliases={'wan':['wan1'],'lan1':['internal','internal1'],'lan2':['wan2'],'lan3':['internal3'],'lan4':['internal4']}
    for key,candidates in aliases.items():
        if result[key] in names:continue
        if current[key]!=MAP[key]:continue
        for candidate in candidates:
            if candidate in names and candidate not in result.values():result[key]=candidate;break
    return result


def verify_changes(s,changes):
    try:
        for op in remediacao.operations(changes):
            value=s[op['category']]['value']
            if op['entry'] is not None:value=next(r for r in value if r['id']==op['entry'])
            actual=value[op['field']]
            if op['field']=='buffer':
                if normalize_banner(actual)!=normalize_banner(op['after']):return False
            elif str(actual)!=op['after']:return False
        return True
    except (KeyError,StopIteration,TypeError):return False

def evaluate(s,p):
    p=effective_policy(s,validate_policy(p)); results=[]
    def add(rule,category,test,recommendation):
        if rule not in p['enabled']: return
        source=s.get(category,{})
        state=UNKNOWN; evidence=source.get('error',s.get('Sistema',{}).get('error','Consulta não disponível')); before=None
        if 'value' in source:
            try: state,evidence,before=test(source['value'])
            except (KeyError,ValueError,TypeError,IndexError): evidence='Campos necessários ausentes ou inválidos'
        item={'rule':rule,'title':LABELS[rule],'status':state,'evidence':evidence,'recommendation':recommendation,'before':before,'after':None,'field':None,'impact':'Requer análise de impacto e procedimento específico.','eligible':False,'blocked':''}
        if rule in AUTO and state=='Não conforme':
            field,limit,_,_=AUTO[rule]; item.update(field=field,after=p[limit],impact=IMPACT[rule],blocked=safety(s));item['eligible']=not item['blocked']
        results.append(item)
    def verdict(ok,text,before=None): return ('Conforme' if ok else 'Não conforme',text,before)
    for rule,(field,limit,_,_) in AUTO.items():
        def numeric(values,r=rule,f=field,k=limit):
            n=int(values[f]); ok=(n>=p[k] if r=='bloqueio' else 0<n<=p[k])
            return verdict(ok,f'{f}: {n}; política: {p[k]}',n)
        add(rule,'Global',numeric,'Ajustar somente este parâmetro ao limite aprovado da política.')
    add('cripto','Global',lambda v:verdict(v['strong-crypto']=='enable','strong-crypto: '+v['strong-crypto']),'Avaliar compatibilidade antes de habilitar strong-crypto.')
    add('tls','Global',lambda v:verdict(bool(v['admin-https-ssl-versions'].split()) and set(v['admin-https-ssl-versions'].split())==set(p['tls_permitidos']),'TLS HTTPS: '+v['admin-https-ssl-versions']+'; padrão esperado: '+' '.join(p['tls_permitidos'])),'Revisar clientes administrativos e permitir somente TLS 1.2/1.3.')
    def interfaces(rows,wan=False):
        if not rows: raise ValueError()
        known=[r for r in rows if 'allowaccess' in r and (not wan or 'role' in r)]
        missing=[r['id'] for r in rows if r not in known]
        bad=[r['id'] for r in known if set(r['allowaccess'].split()) & ({'ssh','https','http','telnet'} if wan else {'http','telnet'}) and (not wan or r['role']=='wan')]
        extra='; campos ausentes nas interfaces: '+', '.join(missing) if missing else ''
        if wan:return ('Revisão necessária' if bad else UNKNOWN if missing else 'Conforme','Interfaces WAN com administração: '+(', '.join(bad) or 'nenhuma confirmada')+extra,None)
        return ('Não conforme' if bad else UNKNOWN if missing else 'Conforme','Interfaces com HTTP/Telnet: '+(', '.join(bad) or 'nenhuma confirmada')+extra,None)
    add('http_telnet','Interfaces',interfaces,'Planejar remoção de HTTP/Telnet preservando o acesso de gerência validado.')
    add('wan_admin','Interfaces',lambda v:interfaces(v,True),'Conferir origem de gerência, trusted hosts e local-in; role WAN sozinho não prova exposição à internet.')
    def trusted(rows):
        if not rows: raise ValueError()
        bad=[]
        for row in rows:
            nets=[ipaddress.IPv4Network(row[f'trusthost{i}'].replace(' ','/'),strict=False) for i in range(1,11)]
            if all(n.prefixlen==0 for n in nets): bad.append(row['id'])
        return verdict(not bad,'Contas sem restrição IPv4 cadastrada: '+(', '.join(bad) or 'nenhuma')+'; IPv6 não avaliado')
    add('trusted_hosts','Administradores',trusted,'Definir origens IPv4/IPv6 aprovadas; validar acesso antes de restringir. A análise atual cobre somente IPv4.')
    def mfa(rows):
        if not rows: raise ValueError()
        locals_=[r for r in rows if r['remote-auth']=='disable']; bad=[r['id'] for r in locals_ if r['two-factor']=='disable']
        return ('Revisão necessária' if bad else 'Conforme','Contas locais sem segundo fator: '+(', '.join(bad) or 'nenhuma')+'; MFA externo e contas de serviço exigem revisão',None)
    add('mfa','Administradores',mfa,'Revisar MFA por conta; distinguir acesso humano, serviço e autenticação externa.')
    def password(values):
        if values['status']=='disable': return verdict(False,'Política local desabilitada')
        if values['status']!='enable': raise ValueError()
        size=int(values['minimum-length']); return verdict(size>=p['senha_min'],f'Política local habilitada; tamanho mínimo: {size}; esperado: {p["senha_min"]}')
    add('senha','Senha',password,'Revisar política de senhas locais e compatibilidade com contas existentes. Não avalia senhas nem política TACACS.')
    add('ntp','NTP',lambda v:verdict(v['ntpsync']=='enable','ntpsync: '+v['ntpsync']+'; estado operacional não verificado'),'Configurar sincronização e verificar alcançabilidade/fontes de horário aprovadas.')
    def snmp(rows):
        enabled=[r for r in rows if r['status']=='enable']
        bad=[r['id'] for r in enabled if any(r[k]=='enable' for k in ('query-v1-status','query-v2c-status','trap-v1-status','trap-v2c-status'))]
        return verdict(not bad,f'Entradas SNMP v1/v2c habilitadas: {len(bad)}; comunidades omitidas')
    add('snmp','SNMP',snmp,'Planejar migração para SNMPv3 e testar o monitoramento antes de desabilitar versões anteriores.')
    for rule,(category,field,target,recommendation) in RULES.items():
        if rule not in p['enabled']:continue
        expected=p['banners']['pre' if rule=='texto_pre' else 'post'] if rule in ('texto_pre','texto_post') else target
        def check(values,r=rule,f=field,wanted=expected):
            if r.startswith('if_'):
                port=p['mapa_interfaces'][r.split('_')[1]]
                matches=[v for v in values if v['id']==port]
                if not matches:return 'Revisão necessária','Interface '+port+' ausente; revise o mapeamento. Não será criada.',None
                if len(matches)!=1:raise ValueError()
                actual=matches[0][f]
            else:actual=values[f]
            if r in ('banner_pre','banner_post'):
                key='pre' if r=='banner_pre' else 'post'
                text=s['BannerPre' if key=='pre' else 'BannerPost']['value']['buffer']
                equal=normalize_banner(text)==normalize_banner(p['banners'][key])
                return verdict(actual==wanted and equal,f'{f}: {actual}; texto institucional: '+('correto' if equal else 'divergente — será incluído na correção'),actual)
            equal=normalize_banner(actual)==normalize_banner(wanted) if r in ('texto_pre','texto_post') else str(actual)==str(wanted)
            return verdict(equal,(f'{r.split("_")[1]} → {p["mapa_interfaces"][r.split("_")[1]]}; ' if r.startswith('if_') else '')+f'{f}: {actual}; esperado: {wanted}',actual)
        if rule=='faz_crypto' and s.get('FortiAnalyzer',{}).get('value',{}).get('status')=='disable':
            results.append({'rule':rule,'title':LABELS[rule],'status':'Não aplicável','evidence':'FortiAnalyzer desativado; não será habilitado sem destino configurado.','recommendation':recommendation,'before':None,'after':None,'field':None,'impact':'Nenhuma alteração','eligible':False,'blocked':''});continue
        if rule=='telemetria_rating' and identity(s)['version'] in SUPPORTED and 'security-rating-result-submission' not in s.get('Global',{}).get('value',{}):
            results.append({'rule':rule,'title':LABELS[rule],'status':'Não aplicável','evidence':'security-rating-result-submission removido no FortiOS 7.4.4. Nesta versão não existe parâmetro para consultar ou aplicar; isso não comprova o estado da telemetria.','recommendation':'Atualizar a referência do script para o firmware em uso. Não enviar comando removido.','before':None,'after':None,'field':None,'impact':'Nenhuma alteração','eligible':False,'blocked':''});continue
        if rule=='maintainer' and identity(s)['version'] in SUPPORTED:
            results.append({'rule':rule,'title':LABELS[rule],'status':'Não aplicável','evidence':'Conta maintainer removida desde FortiOS 7.2.4; comando não será enviado.','recommendation':recommendation,'before':None,'after':None,'field':None,'impact':'Nenhuma alteração','eligible':False,'blocked':''})
            continue
        add(rule,category,check,recommendation)
        item=results[-1]
        if item['status']=='Não conforme':item['after']=expected
        if rule in EXTRA_AUTO and item['status']=='Não conforme':
            item.update(field=field,after=target,impact=recommendation,blocked=safety(s));item['eligible']=not item['blocked']
    for item in results:
        if item['rule'] not in AUTO and item['status'] in ('Não conforme','Revisão necessária'):
            try:
                scoped=copy.deepcopy(s)
                excluded={e['entry'] for e in evolucao.exceptions(p,identity(s)['serial'],item['rule']) if e['entry']}
                if excluded:
                    for cat in ('Administradores','Interfaces','SNMP'):
                        if 'value' in scoped.get(cat,{}):scoped[cat]['value']=[r for r in scoped[cat]['value'] if r['id'] not in excluded]
                ops=remediacao.propose(item['rule'],scoped,p)
                if ops:
                    item.update(operations=ops,before='; '.join(f"{o['entry'] or o['category']}/{o['field']}: {o['before']}" for o in ops),after='; '.join(f"{o['entry'] or o['category']}/{o['field']}: {o['after']}" for o in ops),impact=item['recommendation'],blocked=safety(s))
                    item['eligible']=not item['blocked']
            except (ValueError,KeyError,StopIteration,TypeError) as exc:
                item['eligible']=False
                item['blocked']=str(exc) if isinstance(exc,ValueError) else ('Interface ausente. Use Sugerir mapeamento da última coleta na política e confira os destinos.' if item['rule'].startswith('if_') else 'Campos necessários não coletados; refaça a verificação')
        if 'Session' in s and 'value' not in s['Session'] and item['rule'] in ('wan_admin','http_telnet','trusted_hosts','mfa'):
            item.update(eligible=False,blocked='Identificação da sessão SSH indisponível; corrija a permissão/coleta antes de alterar acesso')
        if item['rule'].startswith('if_'):
            logical=item['rule'].split('_')[1];actual=p['mapa_interfaces'][logical]
            if logical!=actual:item['title']=actual+' — '+RULES[item['rule']][1]+' (script: '+logical+')'
        item['origin']='Script fornecido' if item['rule'] in SCRIPT_IDS else 'Complementar'
    for item in results:
        item['compatibility']=evolucao.describe(s,item['rule'],p)
        if item['status']=='Não aplicável':item['compatibility']='Não aplicável'
    return evolucao.apply_exceptions(results,p,identity(s)['serial'])

def effective_policy(snapshot,policy):
    p=copy.deepcopy(policy);session=snapshot.get('Session',{}).get('value',{})
    if evolucao.model(snapshot)=='60F' or evolucao.legacy_50e(snapshot):p['mapa_interfaces']=suggest_interfaces(snapshot,p['mapa_interfaces'])
    if session:
        p['correcoes']['origem_ssh']=session['source']
        p['correcoes']['interface_gerencia']=session['interface']
    return p


def session_guard(snapshot,changes,policy):
    ops=remediacao.operations(changes)
    risky=any(o['category']=='Administradores' or o['field']=='allowaccess' for o in ops)
    if not risky:return
    session=snapshot.get('Session',{}).get('value',{})
    if not session:raise ValueError('Consulta da sessão SSH atual necessária para alterar acesso/administradores')
    for o in ops:
        if o['category']=='Interfaces' and o['entry']==session['interface'] and o['field']=='allowaccess' and 'ssh' not in o['after'].split():raise ValueError('Plano removeria SSH da interface desta sessão')
        if o['category']=='Administradores' and o['entry']==session['username'] and o['field']=='two-factor':raise ValueError('Use outra conta executora para alterar MFA')
    if any(o['field'].startswith('trusthost') for o in ops):
        nets=[ipaddress.IPv4Network(n) for n in policy['correcoes']['redes_gerencia']]
        if not any(ipaddress.IPv4Address(session['source']) in n for n in nets):raise ValueError('Origem SSH atual fora das redes aprovadas')


def wait_ha(connection,item,seconds,cancel,persist):
    deadline=time.monotonic()+seconds
    while True:
        after=snapshot(connection);reason=safety(after)
        if identity(after)!=item['identity'] or group_id(after)!=item['group']:return after
        if not reason or 'Sincronismo' not in reason:return after
        persist({'event':'ha_wait','device':item['device'],'detail':reason})
        remaining=deadline-time.monotonic()
        if remaining<=0:return after
        if cancel.wait(min(5,remaining)):return after


def collection_error(exc):
    name=type(exc).__name__
    if name in ('AuthenticationException','BadAuthenticationType','PartialAuthentication'):
        return 'Autenticação SSH recusada. Confira usuário/senha, conta permitida e se o acesso exige OTP. Nenhuma configuração foi consultada.'
    if name=='BadHostKeyException':return 'Chave SSH do equipamento diverge da chave confiável; confira a identidade do firewall.'
    if name in ('TimeoutError','socket.timeout'):return 'Tempo de conexão/coleta esgotado. Confira acesso à rede e porta SSH.'
    return 'Falha SSH/coleta: '+name

def audit_device(device,password,timeout,known_hosts,p):
    connection=None
    try:
        connection=SSH(device,password,timeout,known_hosts); s=snapshot(connection)
    except Exception as exc:
        s={'Sistema':{'error':collection_error(exc)}}
    finally:
        if connection: connection.close()
    return {'device':copy.deepcopy(device),'snapshot':s,'identity':identity(s),'findings':evaluate(s,p),'at':datetime.now(timezone.utc).isoformat()}

def audit(devices,password,p,workers=4,timeout=45,known_hosts='',progress=lambda d:None):
    p=validate_policy(p)
    if not 1<=workers<=16: raise ValueError('Concorrência inválida')
    results=[]
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures={pool.submit(audit_device,d,password,timeout,known_hosts,p):d for d in devices}
        for future in as_completed(futures):
            results.append(future.result());progress(futures[future]['nome'])
    return {'policy':p,'devices':sorted(results,key=lambda d:(d['device']['nome'],d['device']['host']))}

def make_plan(report,selections):
    """Recalcula propostas da fotografia auditada; não aceita comandos fornecidos pela interface."""
    plan=[]; seen=set(); selected=set(selections)
    for index,item in enumerate(report['devices']):
        findings=evaluate(item['snapshot'],report['policy'])
        chosen=[]
        for finding in findings:
            if (index,finding['rule']) not in selected: continue
            seen.add((index,finding['rule']))
            if not finding['eligible']: raise ValueError('Item não elegível: '+finding['title'])
            change={k:copy.deepcopy(finding[k]) for k in ('rule','field','before','after','title','impact')}
            if 'operations' in finding:change['operations']=copy.deepcopy(finding['operations'])
            if finding['rule']=='mfa' and any(o['entry']==item['device']['usuario'] for o in change.get('operations',[])):
                raise ValueError('Para habilitar MFA na conta usada pelo programa, conecte com outra conta administrativa. A conta executora precisa continuar apta à verificação SSH sem OTP.')
            if finding['rule']=='wan_admin':
                interface=effective_policy(item['snapshot'],report['policy'])['correcoes']['interface_gerencia']
                port=next(r for r in item['snapshot']['Interfaces']['value'] if r['id']==interface)
                if not item['snapshot'].get('Session',{}).get('value') and item['device']['host']!=port.get('ip','').split(' ')[0]:raise ValueError('Para remover acesso WAN, cadastre o IP direto da interface de gerência preservada; DNS/NAT não comprovam o caminho usado')
            chosen.append(change)
        if chosen:
            remediacao.operations(chosen)
            plan.append({'device':copy.deepcopy(item['device']),'identity':item['identity'].copy(),'group':group_id(item['snapshot']),
                         'snapshot_hash':digest(evolucao.config_state(item['snapshot'])),'interface_mapping':effective_policy(item['snapshot'],report['policy'])['mapa_interfaces'],'changes':chosen})
    if not plan or selected!=seen: raise ValueError('Nenhuma correção aplicável selecionada. Marque apenas itens com aplicação disponível; controles conformes e recomendações manuais não geram comandos.')
    groups=[p['group'] for p in plan]
    if len(set(groups))!=len(groups): raise ValueError('O mesmo equipamento/cluster aparece mais de uma vez; selecione apenas um endereço por cluster')
    return plan

class Writer(SSH):
    def write_global(self,changes):
        # Legacy name retained; the structured plan now includes multiple sections.
        self.write_batches(remediacao.batches(changes))

    def write_batches(self,batches):
        root_prompt=self.prompt
        if '(' in root_prompt or not root_prompt.endswith('#'):raise ValueError('Prompt administrativo não reconhecido')
        name=root_prompt[:-1].strip()
        pattern=re.escape(name)+r'\s+\([^\r\n]+\)\s*#'
        for batch in batches:
            for command in batch:
                self.channel.sendall(command+'\n')
                output=self.receive(prompt_pattern=re.escape(root_prompt) if command=='end' else pattern)
                if ERROR.search(output):raise ValueError('Comando recusado; estado precisa ser conferido')
                self.prompt=output.rstrip().splitlines()[-1].strip()
        self.prompt=root_prompt


def commands(changes):
    return '\n'.join(line for batch in remediacao.batches(changes) for line in batch)


def export(report,folder=None,execution=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    folder=Path(folder or BASE/'resultados');folder.mkdir(parents=True,exist_ok=True)
    path=folder/('hardening_'+datetime.now().strftime('%Y%m%d_%H%M%S_%f')+'.xlsx')
    wb=Workbook();ws=wb.active;ws.title='Hardening'
    ws.append(['Equipamento','Host','Serial','Regra','Resultado','Evidência','Recomendação','Aplicação','Antes','Proposto','Impacto','Origem'])
    summary=wb.create_sheet('Resumo');summary.append(['Equipamento','Host','Conforme','Não conforme','Revisão necessária',UNKNOWN,'Não aplicável','Elegíveis'])
    for device in report['devices']:
        counts={s:0 for s in ('Conforme','Não conforme','Revisão necessária',UNKNOWN,'Não aplicável')}
        for f in device['findings']:
            counts[f['status']]+=1
            ws.append([device['device']['nome'],device['device']['host'],device['identity']['serial'],f['title'],f['status'],f['evidence'],f['recommendation'],application_reason(f),f['before'],f['after'],f['impact'],f.get('origin','')])
        summary.append([device['device']['nome'],device['device']['host'],*counts.values(),sum(f['eligible'] for f in device['findings'])])
    if execution is not None:
        ex=wb.create_sheet('Aplicação');ex.append(['Equipamento','Host','Regra','Antes','Aprovado','Resultado','Detalhe','Backup criptografado','Procedimento'])
        for row in execution:
            for change in row.get('changes',[]):ex.append([row['device']['nome'],row['device']['host'],change['title'],change['before'],change['after'],row['status'],row['detail'],row.get('backup',{}).get('path',''),row.get('backup',{}).get('procedure','')])
    sessions=wb.create_sheet('Sessão SSH');sessions.append(['Equipamento','Usuário','Origem vista pelo firewall','Interface','Destino','Situação'])
    for d in report['devices']:
        session=d['snapshot'].get('Session',{});v=session.get('value',{})
        sessions.append([d['device']['nome'],v.get('username',''),v.get('source',''),v.get('interface',''),v.get('destination',''),session.get('error','Identificada' if v else 'Não coletada')])
    policy_sheet=wb.create_sheet('Política');policy_sheet.append(['Parâmetro','Valor'])
    for key,value in report['policy'].items():policy_sheet.append([key,json.dumps(value,ensure_ascii=False) if isinstance(value,(list,dict)) else value])
    policy_sheet.append(['Gerado em UTC',datetime.now(timezone.utc).isoformat()])
    policy_sheet.append(['Escopo','Hardening técnico inicial. Estado dos campos coletados; não equivale a certificação ou teste funcional de todos os serviços.'])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='203E60');cell.font=Font(color='FFFFFF',bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top')
        for column in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(column)].width=55 if column in (6,7,11) else 28
    compat=wb.create_sheet('Compatibilidade');compat.append(['Equipamento','Modelo','Firmware','Regra','Evidência de suporte'])
    for d in report['devices']:
        for f in d['findings']:compat.append([d['device']['nome'],evolucao.model(d['snapshot']),d['identity']['version'],f['rule'],f.get('compatibility','Não informado')])
    for row in compat:
        for c in row:
            if isinstance(c.value,str):c.data_type='s'
    from refresh_core.storage import write_json
    write_json(path.with_suffix('.json'),report)
    wb.save(path)
    from refresh_core.storage import event
    event('Relatório hardening',path,detail=str(len(report['devices']))+' equipamentos')
    return path

def apply(report,selections,approved_digest,password,timeout=45,known_hosts='',cancel=None,progress=lambda row:None,folder=None,backup_password='',ha_wait=90):
    plan=make_plan(report,selections)
    if digest(plan)!=approved_digest: raise ValueError('O plano mudou; nova aprovação obrigatória')
    cancel=cancel or threading.Event();folder=Path(folder or BASE/'resultados')
    folder.mkdir(parents=True,exist_ok=True)
    journal=folder/('aplicacao_'+datetime.now().strftime('%Y%m%d_%H%M%S_%f')+'.jsonl')
    log=journal.open('x',encoding='utf-8');outcomes=[];stop=False
    def persist(record):
        log.write(json.dumps(record,ensure_ascii=False)+'\n');log.flush()
        import os
        os.fsync(log.fileno())
    try:
        persist({'event':'approval','utc':datetime.now(timezone.utc).isoformat(),'plan_hash':approved_digest,'plan':plan})
        for item in plan:
            row={'device':item['device'],'identity':item['identity'],'group':item['group'],'changes':copy.deepcopy(item['changes']),'status':'Não executado','detail':'Lote interrompido'};connection=None;started=False
            if not stop and not cancel.is_set():
                try:
                    connection=Writer(item['device'],password,timeout,known_hosts)
                    fresh=snapshot(connection)
                    reason=safety(fresh)
                    if reason: raise ValueError(reason)
                    if identity(fresh)!=item['identity'] or group_id(fresh)!=item['group']: raise ValueError('Identidade/cluster mudou')
                    if digest(evolucao.config_state(fresh))!=item['snapshot_hash']: raise ValueError('Configuração/estado mudou desde a auditoria; audite novamente')
                    if cancel.is_set(): raise ValueError('Cancelado antes da alteração')
                    session_guard(fresh,item['changes'],report['policy'])
                    if len(backup_password)<12:raise ValueError('Informe senha do backup com pelo menos 12 caracteres')
                    backup=backup_seguro.save(connection,folder,backup_password,item['identity'])
                    row['backup']=backup
                    persist({'event':'backup','device':item['device'],**backup})
                    # Backup concluído antes de qualquer alteração.
                    persist({'event':'before','device':item['device'],'snapshot':fresh,'changes':item['changes']})
                    started=True
                    connection.write_global(item['changes'])
                    after=wait_ha(connection,item,ha_wait,cancel,persist)
                    reason=safety(after)
                    if reason or identity(after)!=item['identity'] or group_id(after)!=item['group']: raise ValueError('Identidade/HA não confirmado após alteração')
                    expected=remediacao.expected(fresh,item['changes'])
                    if digest(evolucao.config_state(after))!=digest(evolucao.config_state(expected)): raise ValueError('Verificação posterior divergiu; conferir equipamento')
                    persist({'event':'after','device':item['device'],'snapshot':after})
                    row.update(status='Verificado',detail='Campos aprovados conferidos por nova consulta; demais campos auditados preservados.')
                except Exception as exc:
                    stop=True
                    detail=str(exc) if type(exc) is ValueError else type(exc).__name__
                    row.update(status='Estado incerto — revisar' if started else 'Bloqueado',detail=detail+'; sem repetição ou rollback automático.')
                finally:
                    if connection:connection.close()
            persist({'event':'result',**row});outcomes.append(row);progress(row)
    finally:log.close()
    from refresh_core.storage import event
    event('Aplicação hardening',journal,status='Verificada' if all(r['status']=='Verificado' for r in outcomes) else 'Revisar',detail=str(len(outcomes))+' equipamentos; consulte o diário')
    return outcomes,journal
