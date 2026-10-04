"""Localiza MACs em tabelas de switches Cisco/Juniper; consultas SSH somente leitura."""
# Shared core is resolved for both source and portable distributions.
import sys as _sys
from pathlib import Path as _Path
for _parent in _Path(__file__).resolve().parents:
    if (_parent/'refresh_core').is_dir():
        if str(_parent) not in _sys.path:_sys.path.insert(0,str(_parent))
        break
else:raise ImportError('refresh_core ausente: use o pacote completo da Central')

import argparse
import csv
from datetime import datetime, timezone
import getpass
import hashlib
import json
import os
from pathlib import Path
import re
import shlex
from concurrent.futures import ThreadPoolExecutor, as_completed
import xml.etree.ElementTree as ET

BASE=Path(__file__).resolve().parent
MAC_RE=r'(?:[0-9a-fA-F]{2}[:-]){5}[0-9a-fA-F]{2}|(?:[0-9a-fA-F]{4}\.){2}[0-9a-fA-F]{4}'
PORT_RE=r'(?i)\b(?:ge-|xe-|et-|fe-|mge-|ae|Gi(?:gabitEthernet)?|Te(?:nGigabitEthernet)?|Fa(?:stEthernet)?|Eth(?:ernet)?|Po(?:rt-channel)?|Hu(?:ndredGigE)?|Twe(?:ntyFiveGigE)?|Fo(?:rtyGigabitEthernet)?)\d[\d/.:_-]*'


def mac(value):
    value=value.strip()
    if not re.fullmatch(r'(?:[0-9a-fA-F]{12}|'+MAC_RE+r')',value):
        raise ValueError(f'MAC inválido: {value}')
    return re.sub(r'[^0-9a-f]','',value.lower())


def display(value):
    if not value: return ''
    return ':'.join(value[i:i+2] for i in range(0,12,2))


def port_id(value):
    v=value.strip().lower()
    for full,short in [('gigabitethernet','gi'),('tengigabitethernet','te'),('fastethernet','fa'),('ethernet','eth'),('port-channel','po'),('hundredgige','hu'),('twentyfivegige','twe'),('fortygigabitethernet','fo')]:
        if v.startswith(full): v=short+v[len(full):]; break
    if re.match(r'^(ge-|xe-|et-|fe-|mge-|ae\d)',v) and v.endswith('.0'): v=v[:-2]
    return v


def read_csv(path,required,allow_empty=False):
    with Path(path).open(encoding='utf-8-sig',newline='') as f:
        reader=csv.DictReader(f)
        if not set(required)<=set(reader.fieldnames or []): raise ValueError('CSV precisa das colunas: '+','.join(required))
        rows=[]
        for r in reader:
            if None in r: raise ValueError('Colunas extras no CSV; confira o separador vírgula')
            r={k:(v or '').strip() for k,v in r.items()}
            if any(r.values()): rows.append(r)
    if not rows and not allow_empty: raise ValueError('Lista vazia: '+str(path))
    return rows


def parse_fdb(raw,platform):
    rows=[]; skipped=0; instance=''
    for line in raw.splitlines():
        match=re.search(r'Routing instance\s*:\s*(\S+)',line,re.I)
        if match: instance=match[1]
        m=re.search(MAC_RE,line)
        if not m: continue
        before=line[:m.start()].split(); after=line[m.end():]
        ports=re.findall(PORT_RE,after)
        vlan=before[-1] if before else ''
        if not ports:
            skipped+=1; continue
        kind=after.split()[0] if after.split() else ''
        for port in ports:
            rows.append({'mac':mac(m[0]),'vlan':vlan,'instance':instance,'port':port_id(port),'type':kind,'line':line.strip()})
    recognized=bool(rows) or bool(re.search(r'(?i)(mac address table|ethernet switching table|ethernet-switching table|total mac addresses|mac flags|no .*entries)',raw))
    return rows,recognized,skipped


def parse_config(raw,junos):
    ports={}; aliases=[]; current=None
    for line in raw.splitlines():
        if junos:
            try: t=shlex.split(line,comments=True)
            except ValueError: continue
            if len(t)>3 and t[:3]==['set','system','host-name']: aliases.append(t[3])
            if len(t)<4 or t[:2]!=['set','interfaces'] or t[2]=='interface-range': continue
            current=ports.setdefault(port_id(t[2]),{})
            tail=t[3:]
            if tail[:1]==['description']: current['description']=' '.join(tail[1:])
            for key in ('interface-mode','port-mode','802.3ad'):
                if key in tail and tail.index(key)+1<len(tail): current['group' if key=='802.3ad' else 'mode']=tail[tail.index(key)+1]
        else:
            s=line.strip()
            if s.startswith('hostname '): aliases.append(s[9:])
            if s.startswith('interface '): current=ports.setdefault(port_id(s[10:]),{}); continue
            if not line.startswith((' ','\t')): current=None
            if current is None: continue
            if s.startswith('description '): current['description']=s[12:]
            if s.startswith('switchport mode '): current['mode']=s[16:]
            m=re.match(r'channel-group (\d+)',s)
            if m: current['group']='po'+m[1]
    return ports,aliases


def parse_lldp(raw,junos):
    rows=[]
    if junos:
        start=raw.find('<rpc-reply')
        if start<0: start=raw.find('<lldp-neighbors-information')
        if start<0: raise ValueError('Resposta LLDP XML não reconhecida')
        closing='</rpc-reply>' if raw[start:].startswith('<rpc-reply') else '</lldp-neighbors-information>'
        root=ET.fromstring(raw[start:raw.index(closing,start)+len(closing)])
        for node in root.iter(): node.tag=node.tag.split('}')[-1]
        for node in root.iter('lldp-neighbor-information'):
            def text(tag): return (node.findtext('.//'+tag) or '').strip()
            local=text('lldp-local-interface') or text('lldp-local-port-id')
            if local and not local.isdigit():
                rows.append({'port':port_id(local),'neighbor':text('lldp-remote-system-name'),
                             'remote':text('lldp-remote-port-id') or text('lldp-remote-port-description')})
    else:
        current=None
        for line in raw.splitlines():
            m=re.match(r'\s*Local (?:Intf|Interface|Port id)\s*:\s*(\S+)',line,re.I)
            if m: current={'port':port_id(m[1])}; rows.append(current); continue
            if current is None: continue
            for label,key in [('System Name','neighbor'),('Port id','remote')]:
                m=re.match(r'\s*'+label+r'\s*:\s*(.*)',line,re.I)
                if m: current[key]=m[1].strip()
    return rows


def collect(d,password,secret,args,folder):
    from refresh_core.ssh import connect_switch as ConnectHandler
    result={'name':d['nome'],'host':d['host'],'at':datetime.now(timezone.utc).isoformat(),
            'entries':[],'ports':{},'lldp':[],'aliases':[d['nome'],d['host']],'errors':[],'fdb_ok':False}
    junos=d['plataforma']=='juniper_junos'
    def error(cmd,exc):
        msg=str(exc)
        for value in (password,secret):
            if value: msg=msg.replace(value,'[oculto]')
        result['errors'].append({'command':cmd,'error':type(exc).__name__+': '+msg[:1500]})
    try:
        options=dict(device_type=d['plataforma'],host=d['host'],port=int(d.get('porta') or 22),
                     username=d['usuario'],password=password,secret=secret,ssh_strict=True,system_host_keys=True,
                     conn_timeout=args.timeout,auth_timeout=args.timeout,banner_timeout=args.timeout)
        if args.known_hosts: options.update(alt_host_keys=True,alt_key_file=str(args.known_hosts))
        with ConnectHandler(**options) as conn:
            if secret and not junos: conn.enable()
            def query(command,save=True):
                raw=conn.send_command(command,read_timeout=args.timeout)
                if save:
                    fname=hashlib.sha256((d['nome']+command).encode()).hexdigest()[:16]+'.txt'
                    (folder/fname).write_text(d['nome']+'\n'+command+'\n'+raw,encoding='utf-8')
                if re.search(r'(?im)^\s*(?:%\s*(?:Invalid|Error|Ambiguous|Authorization)|error:|syntax error|permission denied|unknown command)|ACCESS-DENIED',raw):
                    raise ValueError('Consulta recusada pelo switch')
                return raw
            commands=['show ethernet-switching table | no-more'] if junos else ['show mac address-table','show mac-address-table']
            for command in commands:
                try:
                    rows,ok,skipped=parse_fdb(query(command),d['plataforma'])
                    if not ok: raise ValueError('Formato da tabela MAC não reconhecido')
                    result['entries']=rows; result['fdb_ok']=True; result['skipped']=skipped
                    if skipped: result['errors'].append({'command':command,'error':f'{skipped} entradas sem porta local reconhecida (CPU, remotas ou formato diferente); consultar log.'})
                    break
                except Exception as exc: error(command,exc)
            command='show configuration | display inheritance | display set | no-more' if junos else 'show running-config'
            try:
                result['ports'],aliases=parse_config(query(command,False),junos)
                result['aliases']+=aliases
            except Exception as exc: error(command,exc)
            command='show lldp neighbors | display xml | no-more' if junos else 'show lldp neighbors detail'
            try: result['lldp']=parse_lldp(query(command),junos)
            except Exception as exc: error(command,exc)
    except Exception as exc: error('SSH',exc)
    return result


def access_locations(rows):
    """One access observation per MAC/switch/port; consolidate VLANs, never uplinks."""
    grouped={}
    for row in rows:
        if row.get('Classificação')!='Candidata de acesso': continue
        key=(row['MAC'],row['Switch'],row['Porta'])
        if key not in grouped:
            grouped[key]=dict(row); grouped[key]['VLAN/nome']=''
        values=set(filter(None,grouped[key]['VLAN/nome'].split(', ')))
        values.add(str(row['VLAN/nome']))
        grouped[key]['VLAN/nome']=', '.join(sorted(values))
    return list(grouped.values())


def analyze(targets,results):
    aliases={}
    for device in results:
        for alias in device['aliases']:
            keys={alias.lower()}
            if not re.fullmatch(r'[\d.]+',alias): keys.add(alias.lower().split('.')[0])
            for key in keys: aliases.setdefault(key,set()).add(device['name'])
    by_mac={}
    for device in results:
        for entry in device['entries']:
            by_mac.setdefault(entry['mac'],[]).append((device,entry))
    hits=[]; summary=[]
    incomplete=[d['name'] for d in results if not d['fdb_ok'] or d.get('skipped',0) or d['errors']]
    seen_targets=set()
    for target in targets:
        identity=target['mac'] or (target.get('ip',''),target.get('interface_firewall',''))
        if identity in seen_targets: continue
        seen_targets.add(identity)
        found=[]
        for device,entry in by_mac.get(target['mac'],[]):
            port=entry['port']; cfg=device['ports'].get(port,{})
            members=[p for p,c in device['ports'].items() if port_id(c.get('group',''))==port]
            neighbors=[n for n in device['lldp'] if n['port'] in [port]+members]
            downstream=set()
            for n in neighbors:
                value=n.get('neighbor','').lower()
                downstream |= aliases.get(value,set()) or aliases.get(value.split('.')[0],set())
            others=downstream-{device['name']}
            kind='Candidata de acesso' if cfg.get('mode')=='access' and not others and not members else 'Porta de aprendizado; validar conexão final'
            if others: kind='Enlace para switch cadastrado'
            elif members or cfg.get('group') or re.match(r'^(ae\d|po\d)',port): kind='Agregação; membro físico não determinado'
            elif cfg.get('mode')=='trunk': kind='Trunk; dispositivo pode estar adiante'
            if entry['type'].lower() not in ('dynamic','d','learn','learned'):
                kind='Entrada estática/especial; não comprova presença'
            row={'Dispositivo':target['nome'],'MAC':display(target['mac']),'Switch':device['name'],'IP switch':device['host'],
                 'Porta':port,'VLAN/nome':entry['vlan'],'Instância':entry['instance'],'Tipo MAC':entry['type'],
                 'Descrição':cfg.get('description',''),'Modo':cfg.get('mode','não extraído'),'Classificação':kind,
                 'LLDP':'; '.join(n.get('neighbor','')+' / '+n.get('remote','') for n in neighbors),
                 'Próximos switches':', '.join(sorted(others)),'Membros agregado':', '.join(members),
                 'Coleta UTC':device['at'],'Evidência':entry['line']}
            found.append(row)
        for row in found:
            next_names=row['Próximos switches'].split(', ') if row['Próximos switches'] else []
            next_hits=[r for r in found if r['Switch'] in next_names]
            row['MAC nos próximos switches']='; '.join(f"{r['Switch']} / {r['Porta']} / VLAN {r['VLAN/nome']}" for r in next_hits)
            if next_names and not next_hits: row['MAC nos próximos switches']='Não observado; conferir cobertura e VLAN'
        access=access_locations(found)
        sources=target.get('origem','').split('; ')
        if not access and all(source.startswith('Tabela MAC / ') for source in sources): continue
        hits.extend(access)
        candidates=[f"{r['Switch']} / {r['Porta']} / VLAN {r['VLAN/nome']}" for r in access]
        status='Não encontrado nesta coleta' if not found else 'Uma porta candidata de acesso' if len(candidates)==1 else 'Múltiplas candidatas; revisar' if candidates else 'Aprendido em enlace/agregação ou modo não identificado'
        if not target['mac']: status='MAC não identificado; porta não determinada'
        if incomplete: status+=' — cobertura parcial'
        if not target['mac'] or len(candidates)!=1 or incomplete or target.get('revisar'): status='A revisar — '+status
        summary.append({'Dispositivo':target['nome'],'MAC':display(target['mac']),'IP dispositivo':target.get('ip',''),'Resultado':status,
                        'Origem':target.get('origem','Lista manual'),'Interface firewall':target.get('interface_firewall',''),
                        'Candidatas':'; '.join(candidates),'Ocorrências':len(found),'Switches com cobertura parcial/falha':', '.join(incomplete),
                        'Observação':'Tabela MAC indica aprendizado, não garante conexão direta. AP, telefone, hipervisor ou switch intermediário podem estar na porta. '+target.get('revisar','')})
    return {'Localização':summary,'Ocorrências':hits,'Coleta':[{'Switch':d['name'],'IP':d['host'],'Tabela MAC': 'OK' if d['fdb_ok'] else 'FALHA',
            'Entradas':len(d['entries']),'Coleta UTC':d['at'],'Erros':'; '.join(e['command']+': '+e['error'] for e in d['errors'])} for d in results]}


def excel(tables,path):
    from modelo_excel import prepare_template
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    w=prepare_template(tables)
    for name,rows in tables.items():
        s=w.create_sheet(name); headers=list(dict.fromkeys(k for r in rows for k in r)) or ['Informação']
        s.append(headers)
        if not rows: s.append(['Sem ocorrências; confira a cobertura na aba Coleta.'])
        for r in rows: s.append([r.get(k,'') for k in headers])
        for row in s:
            for c in row:
                if isinstance(c.value,str): c.value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',c.value)[:32767]; c.data_type='s'
                c.alignment=Alignment(vertical='top',wrap_text=True)
        for c in s[1]: c.font=Font(color='FFFFFF',bold=True); c.fill=PatternFill('solid',fgColor='17365D')
        for i in range(1,len(headers)+1): s.column_dimensions[get_column_letter(i)].width=30
        s.freeze_panes='C2'; s.auto_filter.ref=s.dimensions
    with path.open('xb') as f: w.save(f)


def run_collection(jobs,fw_jobs,targets,args,output_dir=None,on_event=None):
    import fortigate
    completed=0; total=len(jobs)+len(fw_jobs)
    def notify(message):
        if on_event: on_event({'message':message,'completed':completed,'total':total})
        else: print(message)
    notify('Iniciando consultas SSH...')
    run=Path(output_dir or BASE/'resultados')/datetime.now().strftime('%Y%m%d_%H%M%S_%f'); logs=run/'logs'; logs.mkdir(parents=True)
    results=[]; fw_results=[]
    # Query each configured cluster once; switches still run concurrently below.
    for d,pw in fw_jobs:
        notify('Consultando FortiGate '+d['nome']+'…')
        r=fortigate.collect(d,pw,args,logs); fw_results.append(r)
        completed+=1
        notify(r['name']+': '+str(len(r['hosts']))+' registros ARP/DHCP; '+('A revisar' if r['errors'] else 'coletado'))
        for e in r['errors']: notify('  '+e['command']+': '+e['error'])
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        notify('Consultando '+str(len(jobs))+' switches…')
        futures={pool.submit(collect,d,pw,sec,args,logs):d for d,pw,sec in jobs}
        for future in as_completed(futures):
            d=futures[future]
            try: r=future.result()
            except Exception as exc: r={'name':d['nome'],'host':d['host'],'at':'','entries':[],'ports':{},'lldp':[],'aliases':[d['nome']],'fdb_ok':False,'errors':[{'command':'Coletor','error':type(exc).__name__}]}
            results.append(r); completed+=1; notify(r['name']+': '+('tabela MAC coletada' if r['fdb_ok'] else 'FALHA'))
            for e in r['errors']: notify('  '+e['command']+': '+e['error'])
    results.sort(key=lambda r:r['name'])
    (run/'coleta.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    if args.modo=='integrado':
        targets=fortigate.discover(fw_results,results,targets)
        (run/'fortigate.json').write_text(json.dumps(fw_results,ensure_ascii=False,indent=2),encoding='utf-8')
    tables=analyze(targets,results)
    if args.modo=='integrado': fortigate.add_tables(tables,fw_results)
    else: tables['A revisar']=[r.copy() for r in tables['Localização'] if r['Resultado'].startswith('A revisar')]
    notify('Gerando o Excel e a lista de itens a revisar…')
    path=run/'localizacao_macs.xlsx'; excel(tables,path); notify('Excel: '+str(path))
    notify(str(len(tables['Localização']))+' registros; '+str(len(tables['A revisar']))+' a revisar.')
    from refresh_core.storage import event
    event('Localização MAC',path,detail=str(len(tables['Localização']))+' registros')
    return {'path':path,'tables':tables,'code':2 if any(d['errors'] for d in results+fw_results) else 0}


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--switches',type=Path,default=BASE/'switches.csv'); p.add_argument('--macs',type=Path,help='Lista opcional adicional; no modo lista usa macs.csv por padrão')
    p.add_argument('--firewalls',type=Path,default=BASE/'firewalls.csv')
    p.add_argument('--modo',choices=('integrado','lista'),default='integrado')
    p.add_argument('--usuario',default=os.getenv('SW_USER','')); p.add_argument('--known-hosts',type=Path)
    p.add_argument('--timeout',type=int,default=90); p.add_argument('--workers',type=int,default=4)
    args=p.parse_args()
    try:
        import netmiko,openpyxl
        import fortigate
        devices=read_csv(args.switches,['nome','host','plataforma'])
        targets=read_csv(args.macs or BASE/'macs.csv',['nome','mac']) if args.macs or args.modo=='lista' else []
        firewalls=[]
        if args.modo=='integrado':
            firewalls=read_csv(args.firewalls,['nome','host'],allow_empty=True) if args.firewalls.exists() else []
            if not firewalls:
                print('Cadastre o endereço de gerenciamento do cluster FortiGate. Será salvo sem senha em '+str(args.firewalls))
                host=input('IP ou DNS do FortiGate: ').strip()
                name=input('Nome do FortiGate: ').strip()
                if not host or not name: raise ValueError('Informe nome e endereço do FortiGate')
                firewalls=[dict(nome=name,host=host,porta='22',usuario='',password_env='FG_PASSWORD')]
                with args.firewalls.open('w',encoding='utf-8-sig',newline='') as f:
                    writer=csv.DictWriter(f,fieldnames=list(firewalls[0])); writer.writeheader(); writer.writerows(firewalls)
        if not 1<=args.workers<=16 or not 5<=args.timeout<=600: raise ValueError('workers: 1..16; timeout: 5..600')
        for t in targets: t['mac']=mac(t['mac'])
        seen=set(); endpoints=set()
        for d in devices:
            port=int(d.get('porta') or 22)
            if not d['nome'] or not d['host'] or d['plataforma'] not in ('cisco_ios','cisco_nxos','juniper_junos') or not 1<=port<=65535: raise ValueError('Switch com campos inválidos')
            if d['nome'].lower() in seen or (d['host'].lower(),port) in endpoints: raise ValueError('Switch duplicado na lista')
            seen.add(d['nome'].lower()); endpoints.add((d['host'].lower(),port))
        fw_names=set(); fw_endpoints=set()
        for d in firewalls:
            port=int(d.get('porta') or 22)
            if not d['nome'] or not d['host'] or not 1<=port<=65535: raise ValueError('FortiGate com nome, host ou porta inválidos')
            if d['nome'].lower() in fw_names or (d['host'].lower(),port) in fw_endpoints: raise ValueError('FortiGate duplicado na lista')
            fw_names.add(d['nome'].lower()); fw_endpoints.add((d['host'].lower(),port))
        secrets={}; jobs=[]; fw_jobs=[]
        for d,is_fw in [(d,False) for d in devices]+[(d,True) for d in firewalls]:
            d['usuario']=d.get('usuario') or args.usuario
            if not d['usuario']: args.usuario=d['usuario']=input('Usuário SSH ('+d['nome']+'): ').strip()
            if not d['usuario']: raise ValueError('Usuário vazio')
            def credential(env):
                if not env: return ''
                if env in os.environ: return os.environ[env]
                if env not in secrets: secrets[env]=getpass.getpass('Credencial '+env+' (não será salva): ')
                return secrets[env]
            if is_fw: fw_jobs.append((d,credential(d.get('password_env') or 'FG_PASSWORD')))
            else: jobs.append((d,credential(d.get('password_env') or 'SW_PASSWORD'),credential(d.get('secret_env',''))))
        return run_collection(jobs,fw_jobs,targets,args)['code']
    except (ValueError,OSError,ImportError) as exc: print('Erro: '+str(exc)); return 1


if __name__=='__main__':
    import sys
    if len(sys.argv)==1:
        from interface import main as gui_main
        gui_main()
    else: raise SystemExit(main())
