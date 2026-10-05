"""Adaptação ao modelo Verificacao_Portas_C001; não altera o original."""
from copy import copy
from pathlib import Path
import re


def interface_id(value):
    value = str(value or '').strip().lower()
    for long, short in [('hundredgigabitethernet', 'hu'), ('fortygigabitethernet', 'fo'),
                        ('twentyfivegige', 'twe'), ('tengigabitethernet', 'te'),
                        ('gigabitethernet', 'gi'), ('fastethernet', 'fa'),
                        ('port-channel', 'po'), ('ethernet', 'eth')]:
        if value.startswith(long):
            value = short + value[len(long):]
            break
    return value


def first(row, *keys):
    for key in keys:
        value = row.get(key)
        if value is not None and value != '' and value != []:
            return ','.join(map(str, value)) if isinstance(value, list) else str(value)
    return ''


def lldp_lines(records):
    """Consolidate summary/XML/detail observations without merging distinct ports."""
    grouped=[]
    for row in records:
        neighbor=first(row,'neighbor_name','neighbor','destination_host','system_name')
        remote=first(row,'neighbor_interface','remote_port','port_id')
        description=first(row,'port_description','neighbor_port_description')
        chassis=first(row,'chassis_id','neighbor_id')
        # Junos may advertise a numeric port ID while the description is the interface.
        if (not remote or remote.isdigit()) and re.match(r'(?i)^(?:ge-|xe-|et-|fe-|ae|gi|te|fa|eth|po)\d',description):remote=description
        canonical=interface_id(remote).removesuffix('.0')
        identity=re.sub(r'[:.\-]','',chassis.lower())
        match=next((x for x in grouped if x['port']==canonical and
                    ((identity and x['identity']==identity) or (not identity or not x['identity']) and neighbor and x['neighbor'].casefold()==neighbor.casefold())),None)
        if match is None:
            match={'port':canonical,'identity':identity,'neighbor':neighbor,'remote':remote,'description':description,'chassis':chassis};grouped.append(match)
        else:
            for key,value in [('neighbor',neighbor),('remote',remote),('description',description),('chassis',chassis),('identity',identity)]:
                if not match[key]:match[key]=value
    lines=[]
    for item in grouped:
        values=[]
        for key in ('neighbor','remote','description','chassis'):
            value=item[key]
            if value and value.casefold() not in {v.casefold() for v in values}:values.append(value)
        line=' | '.join(values)
        if line and line not in lines:lines.append(line)
    return lines


def build_rows(tables):
    ports = {}
    def get(device, name):
        if re.match(r'^(?:ge-|xe-|et-|fe-|ae\d)', name) and name.endswith('.0'):
            name = name[:-2]
        return ports.setdefault((device, interface_id(name)), {'switch': device, 'porta': name})
    mapping = {'descricao': 'descricao', 'modo_configurado': 'modo', 'vlan_access': 'access',
               'vlans_trunk': 'trunk', 'vlans_membros': 'trunk', 'vlan_nativa': 'nativa',
               'agregacao': 'agregado', 'admin_configurado': 'admin'}
    for row in tables.get('Interfaces_config', []):
        name = row['interface']
        # Ethernet-switching unit zero belongs to the physical/aggregate port in this layout.
        if re.match(r'^(?:ge-|xe-|et-|ae\d)', name) and name.endswith('.0'):
            name = name[:-2]
        port = get(row['equipamento'], name)
        for source, target in mapping.items():
            if row.get(source) not in (None, ''):
                port[target] = str(row[source])
    fields = {
        'Estado_portas': {'link': ('status', 'link_status', 'link'), 'admin': ('admin_status', 'admin'),
                         'speed': ('speed',), 'duplex': ('duplex',), 'midia': ('type', 'media_type')},
        'Detalhes_portas': {'link': ('link_status',), 'admin': ('admin_state', 'admin_status'),
                            'speed': ('speed',), 'duplex': ('duplex', 'link_mode'),
                            'midia': ('media_type',), 'rx': ('input_errors',), 'tx': ('output_errors',)},
        'Switchport': {'modo': ('mode', 'admin_mode', 'administrative_mode'),
                       'access': ('access_vlan',), 'trunk': ('trunking_vlans', 'trunk_vlans'),
                       'nativa': ('native_vlan',)},
        'PoE': {'poe': ('status', 'oper', 'oper_status')},
    }
    for sheet, mapping in fields.items():
        for row in tables.get(sheet, []):
            name = first(row, 'interface', 'port', 'interface_name')
            if not name:
                continue
            # Do not let an inet logical-unit state override the physical link.
            if '.' in name and name.startswith(('ge-', 'xe-', 'et-', 'ae')):
                continue
            port = get(row['equipamento'], name)
            for target, aliases in mapping.items():
                value = first(row, *aliases)
                if value:
                    if target == 'link' and port.get('link') in ('err-disabled', 'disabled', 'notconnect'):
                        continue
                    port[target] = value
    for row in tables.get('STP', []):
        name = first(row, 'interface', 'port')
        if name:
            port = get(row['equipamento'], name)
            instance = first(row, 'vlan_id', 'vlan', 'instance')
            for target, aliases in [('stp_role', ('role',)), ('stp_state', ('status', 'state'))]:
                value = first(row, *aliases)
                if value:
                    entry = f'{instance}: {value}' if instance else value
                    port[target] = '; '.join(filter(None, [port.get(target), entry]))
    for sheet in ('Vizinhos_LLDP', 'Vizinhos_CDP'):
        for row in tables.get(sheet, []):
            name = first(row, 'local_interface', 'local_port', 'local_port_id')
            if name:
                port = get(row['equipamento'], name)
                neighbor = first(row, 'neighbor_name', 'neighbor', 'destination_host', 'system_name')
                remote = first(row, 'neighbor_interface', 'remote_port', 'port_id')
                if sheet=='Vizinhos_CDP' and (neighbor or remote):
                    port['obs'] = '; '.join(filter(None, [port.get('obs'), f'{neighbor} {remote}'.strip()]))
                if sheet == 'Vizinhos_LLDP':
                    port.setdefault('lldp_records', []).append(row)
    for sheet in ('LAG_operacional', 'LACP_detalhes'):
        for row in tables.get(sheet, []):
            name = first(row, 'interface')
            if not name:
                continue
            port = get(row['equipamento'], name)
            for source, target in [('agregacao', 'agregado'), ('role', 'lacp_role'), ('state', 'lacp_state'),
                                    ('priority', 'lacp_priority'), ('system_id', 'lacp_sysid')]:
                if row.get(source):
                    port[target] = row[source]
    for row in tables.get('Agregacoes_config', []):
        port = get(row['equipamento'], row['membro'])
        port['agregado'] = row['agregacao']
        port.setdefault('lacp_role', row.get('modo', ''))
    lacp = []
    for port in ports.values():
        if not port.get('agregado'):
            continue
        agg = ports.get((port['switch'], interface_id(port['agregado'])), {})
        lacp.append([port['agregado'], port['switch'], port['porta'], port.get('lacp_role', ''),
                     port.get('lacp_state', 'Não coletado'), port.get('lacp_priority', 'Não coletado'), port.get('lacp_sysid', 'Não coletado'), port.get('speed'), port.get('duplex'),
                     agg.get('trunk'), agg.get('nativa'), agg.get('link'),
                     ''])
    statuses = {'up': '🟢 Up', 'connected': '🟢 Up', 'down': '🔴 Down',
                'notconnect': '🟡 Not connect', 'disabled': '⬜ Disabled', 'err-disabled': '🔴 Err-disabled'}
    rows = []
    def natural(item):
        return [int(x) if x.isdigit() else x for x in re.split(r'(\d+)', str(item[0]))]
    for _, port in sorted(ports.items(), key=natural):
        link = port.get('link', '')
        obs = port.get('obs', '')
        if port.get('agregado'):
            obs += '; VLANs do agregado: consultar a linha ' + port['agregado']
        missing = [label for key, label in [('link', 'link'), ('speed', 'velocidade'),
                    ('stp_role', 'STP'), ('poe', 'PoE'), ('rx', 'erros RX'), ('tx', 'erros TX')]
                   if not port.get(key)]
        if missing:
            obs += '; Não coletado: ' + ', '.join(missing)
        rows.append([port['switch'], port['porta'], port.get('descricao', ''), statuses.get(link.lower(), link) or 'Não coletado',
                     port.get('admin', 'Não coletado'), port.get('speed', 'Não coletado'),
                     port.get('duplex', 'Não coletado'), port.get('modo', ''), port.get('access', ''),
                     port.get('trunk', ''), port.get('nativa', ''), port.get('agregado', ''),
                     port.get('lacp_role', ''), port.get('lacp_state', 'Não coletado') if port.get('agregado') else '',
                     port.get('stp_role', 'Não coletado'), port.get('stp_state', 'Não coletado'),
                     port.get('poe', 'Não coletado'), port.get('midia', 'Não coletado'),
                     f"{port.get('rx', '?')} / {port.get('tx', '?')}", obs.strip('; '),
                     '\n'.join(lldp_lines(port.get('lldp_records', []))) or 'Não informado nesta coleta'])
    return rows, lacp


def vlan_rows(tables):
    result = {}
    for sheet in ('VLANs_operacionais', 'VLANs_config'):
        for row in tables.get(sheet, []):
            device = row.get('equipamento', '')
            vid = first(row, 'vlan_id', 'vlan_tag', 'vlan-id-list')
            name = first(row, 'nome', 'vlan_name', 'name')
            if not vid and not name:
                continue
            instance = first(row, 'instance', 'routing_instance')
            if instance in ('default-switch', 'default'):
                instance = ''
            key = (device, instance, vid or name)
            target = result.setdefault(key, [device, vid, name, instance, '', '', '', '', '', ''])
            for col, value in [(2, name), (4, first(row, 'descricao', 'description')),
                               (5, first(row, 'status', 'vlan_status')), (6, first(row, 'interfaces')),
                               (7, first(row, 'l3-interface')), (9, first(row, 'coleta_utc'))]:
                if value:
                    target[col] = value
            source = first(row, 'fonte') or sheet
            target[8] = '; '.join(filter(None, [target[8], source]))
    return sorted(result.values(), key=lambda r: (r[0], r[3], int(r[1]) if str(r[1]).isdigit() else 4095, str(r[1])))


def load_model(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for sheet, headers in {'🔌 Portas': ['Switch', 'Porta', 'Descrição'],
                           '⚡ LACP': ['Port-Channel', 'Switch', 'Interface Membro']}.items():
        if sheet not in wb.sheetnames or [wb[sheet].cell(2, c).value for c in range(1, 4)] != headers:
            raise ValueError(f'Modelo incompatível: cabeçalhos da aba {sheet}')
    return wb


def export_model(tables, template, output):
    from openpyxl.utils import get_column_letter
    template, output = Path(template), Path(output)
    if template.resolve() == output.resolve():
        raise ValueError('O destino não pode ser o modelo original')
    wb = load_model(template)
    port_rows, lacp_rows = build_rows(tables)
    ports_sheet = wb['🔌 Portas']
    ports_sheet['U2'] = 'LLDP — Vizinho | Porta | Descrição | Chassis'
    ports_sheet['U2']._style = copy(ports_sheet['T2']._style)
    ports_sheet['U3']._style = copy(ports_sheet['T3']._style)
    ports_sheet.column_dimensions['U'].width = 65
    if 'A1:T1' in ports_sheet.merged_cells:
        ports_sheet.unmerge_cells('A1:T1')
        ports_sheet.merge_cells('A1:U1')
    for name, rows, width in [('🔌 Portas', port_rows, 21), ('⚡ LACP', lacp_rows, 13)]:
        ws = wb[name]
        styles = [copy(ws.cell(3, c)._style) for c in range(1, width + 1)]
        height = ws.row_dimensions[3].height
        # Remove historical values even when the new collection failed or has fewer ports.
        for row in ws.iter_rows(min_row=3, max_row=max(3, ws.max_row), max_col=width):
            for cell in row:
                cell.value = None
                cell.hyperlink = None
                cell.comment = None
        if not rows:
            rows = [['Sem dados nesta coleta; consulte a aba Coleta.']]
        for r, values in enumerate(rows, 3):
            ws.row_dimensions[r].height = height
            for c, value in enumerate(values, 1):
                cell = ws.cell(r, c)
                cell._style = copy(styles[c - 1])
                cell.value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))[:32767] if value is not None else ''
                cell.data_type = 's'
        end = len(rows) + 2
        ws.auto_filter.ref = f'A2:{get_column_letter(width)}{end}'
        ws.print_area = f'A1:{get_column_letter(width)}{end}'
        for table in ws.tables.values():
            table.ref = f'A2:{get_column_letter(width)}{end}'
    output.parent.mkdir(parents=True, exist_ok=True)
    if 'VLANs' in wb.sheetnames:
        del wb['VLANs']
    vlans = wb.create_sheet('VLANs')
    vlans.append(['Switch', 'VLAN ID', 'Nome', 'Instância', 'Descrição', 'Estado', 'Interfaces', 'Interface L3', 'Fonte', 'Coleta UTC'])
    for values in vlan_rows(tables):
        vlans.append(values)
    if vlans.max_row == 1:
        vlans.append(['Sem VLANs extraídas; consulte a aba Coleta.'])
    for cells in vlans:
        for cell in cells:
            if isinstance(cell.value, str):
                cell.value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', cell.value)[:32767]
                cell.data_type = 's'
    for cell in vlans[1]:
        cell._style = copy(ports_sheet['A2']._style)
    vlans.freeze_panes = 'C2'
    vlans.auto_filter.ref = vlans.dimensions
    for i in range(1, 11):
        vlans.column_dimensions[get_column_letter(i)].width = 45 if i in (5, 7, 9) else 25
    if 'Coleta' in wb.sheetnames:
        del wb['Coleta']
    audit = wb.create_sheet('Coleta')
    audit.append(['Switch', 'Host', 'Coleta UTC', 'Status / comando', 'Ocorrência'])
    for row in tables.get('Resumo', []):
        audit.append([row.get('equipamento'), row.get('host'), row.get('coleta_utc'), row.get('status'), ''])
    for row in tables.get('Ocorrencias', []):
        audit.append([row.get('equipamento'), row.get('host'), row.get('coleta_utc'), row.get('comando'), row.get('erro')])
    for cells in audit:
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    audit.freeze_panes = 'A2'
    audit.auto_filter.ref = audit.dimensions
    for col, width in [('A', 25), ('B', 20), ('C', 30), ('D', 55), ('E', 90)]:
        audit.column_dimensions[col].width = width
    temp = output.with_suffix('.tmp.xlsx')
    wb.save(temp)
    temp.replace(output)
