"""Preenche o modelo incorporado sem escolher uma localização ambígua."""
from pathlib import Path
from copy import copy
import re
from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties


def text_cell(cell, value):
    cell.value=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value or ''))[:32767]
    cell.data_type='s'


def prepare_template(tables):
    book=load_workbook(Path(__file__).with_name('modelo_correspondencia.xlsx'))
    sheet=book['Mapeamento de Portas']
    styles=[copy(c._style) for c in sheet[4]]
    for row in sheet.iter_rows(min_row=4):
        for cell in row: cell.value=None
    sheet['A2']='Somente portas de acesso, sem uplinks ou repetições. Pendências: A revisar. Detalhes nas abas Localização e Coleta.'
    for column,width in zip('ABCDEFG',[22,29,23,25,19,22,55]): sheet.column_dimensions[column].width=width
    hits={}
    from localizar import access_locations
    for hit in access_locations(tables['Ocorrências']): hits.setdefault(hit['MAC'],[]).append(hit)
    for number,item in enumerate(tables['Localização'],4):
        found=hits.get(item['MAC'],[])
        locations=list(dict.fromkeys((r['Switch'],r['Porta'],r['VLAN/nome']) for r in found))
        columns=list(locations[0]) if len(locations)==1 else ['A revisar','A revisar','']
        if not item['MAC']: notes='A revisar: MAC não identificado.'
        elif not locations: notes='A revisar: porta de acesso não identificada.'
        elif len(locations)>1: notes=f'A revisar: {len(locations)} portas de acesso candidatas. Veja Ocorrências.'
        elif item.get('Resultado','').startswith('A revisar'): notes='Porta de acesso identificada; coleta parcial ou conflito. Veja Localização e Coleta.'
        else: notes='Porta de acesso identificada.'
        ips=item.get('IP dispositivo','').splitlines()
        ip=ips[0] if len(ips)<=1 and ips else ''
        if len(ips)>1: ip=f'{len(ips)} IPs; veja Localização'
        values=[item['MAC'] or 'A revisar',item['Dispositivo'],ip,*columns,notes]
        for col,value in enumerate(values,1):
            cell=sheet.cell(number,col); cell._style=copy(styles[col-1]); text_cell(cell,value)
            alignment=copy(cell.alignment); alignment.wrap_text=True; alignment.vertical='top'; cell.alignment=alignment
        sheet.row_dimensions[number].height=45
    last=max(4,len(tables['Localização'])+3)
    for table in sheet.tables.values():
        table.ref=f'A3:G{last}'
        if table.autoFilter: table.autoFilter.ref=table.ref
    sheet.auto_filter.ref=f'A3:G{last}'; sheet.freeze_panes='D4'
    search=book['Busca por MAC']; search['B4']=None; search['B4'].number_format='@'
    # Normalize the input once in a hidden cell; MATCH uses plain ranges, no array formula.
    search['D4']='=SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE("M"&UPPER(B4),":",""),"-",""),".","")," ","")'
    search['D5']='='+'&":"&'.join(f'MID(D4,{i},2)' for i in range(2,13,2))
    search['D6']=f'=IF(LEN(D4)<>13,0,IFERROR(MATCH(D5,\'Mapeamento de Portas\'!$A$4:$A${last},0),0))'
    search.column_dimensions['D'].hidden=True
    for row,col in enumerate('BCDEFG',6):
        lookup=f'INDEX(\'Mapeamento de Portas\'!${col}$4:${col}${last},$D$6)'
        search.cell(row,2,f'=IF($B$4="","",IF($D$6=0,"MAC não encontrado",IF({lookup}="","",{lookup})))')
        alignment=copy(search.cell(row,2).alignment); alignment.wrap_text=True; search.cell(row,2).alignment=alignment
        search.row_dimensions[row].height=60 if row==11 else 32
        for colnum in (1,2):
            alignment=copy(search.cell(row,colnum).alignment); alignment.vertical='top'; search.cell(row,colnum).alignment=alignment
    search['A13']='Somente portas de acesso. Se houver ambiguidade, consulte Ocorrências e Localização.'
    alignment=copy(search['A13'].alignment); alignment.wrap_text=True; search['A13'].alignment=alignment
    search.row_dimensions[13].height=36
    instructions=book['Instruções']
    instructions['B3']='Edite firewalls.csv e switches.csv. Execute localizar.py para descobrir dispositivos por SSH; lista de MACs é opcional.'
    instructions['B4']='A busca aceita MAC com dois-pontos, hífens, pontos ou sem separadores.'
    instructions['B5']='Na aba Busca por MAC, digite o MAC no campo amarelo. O Excel calcula os resultados ao abrir.'
    instructions['B6']='Uplinks não são exibidos. Portas de acesso repetidas são consolidadas. Ambiguidades ficam A revisar; consulte Ocorrências.'
    instructions['B7']='IP e nome vêm de ARP/DHCP ou lista manual. A revisar indica informação faltante ou ambígua. Lease e ARP não comprovam presença atual.'
    book.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True,forceFullCalc=True)
    book.active=book.sheetnames.index('Mapeamento de Portas')
    return book
