"""Uma sessão SSH por switch; consultas comuns reutilizadas em memória."""
from pathlib import Path
from datetime import datetime,timezone
from concurrent.futures import ThreadPoolExecutor,as_completed
import json
import re
import localizar
import mapear
import fortigate
from refresh_core.ssh import connect_switch

class CachedConnection:
    def __init__(self,connection):self.connection=connection;self.cache={};self.enabled=False;self.failure=None
    def send_command(self,command,**kwargs):
        if self.failure is not None:raise self.failure
        if command not in self.cache:
            try:self.cache[command]=self.connection.send_command(command,**kwargs)
            except Exception as exc:
                self.failure=exc;raise
        return self.cache[command]
    def enable(self):
        if not self.enabled:self.connection.enable();self.enabled=True

def failed(d):
    return {'name':d['nome'],'host':d['host'],'at':datetime.now(timezone.utc).isoformat(),'entries':[],'ports':{},'lldp':[],
            'aliases':[d['nome'],d['host']],'errors':[],'fdb_ok':False}

def collect_switch(d,pw,secret,args,folder):
    meta={'equipamento':d['nome'],'host':d['host'],'coleta_utc':datetime.now(timezone.utc).isoformat()}
    ports={};location=failed(d)
    options=dict(device_type=d['plataforma'],host=d['host'],port=int(d.get('porta') or 22),username=d['usuario'],password=pw,secret=secret,
                 ssh_strict=True,system_host_keys=True,conn_timeout=args.timeout,auth_timeout=args.timeout,banner_timeout=args.timeout)
    if args.known_hosts:options.update(alt_host_keys=True,alt_key_file=str(args.known_hosts))
    if d.get('key_file'):options.update(use_keys=True,key_file=d['key_file'])
    credentials={'username':d['usuario'],'password':pw,'secret':secret}
    try:
        with connect_switch(**options) as raw:
            connection=CachedConnection(raw)
            if not args.rapido:
                try:ports=mapear.collect(d,credentials,args.timeout,args.known_hosts,connection=connection)
                except Exception as exc:ports={'Resumo':[{**meta,'status':'FALHA'}],'Ocorrencias':[{**meta,'comando':'Portas','erro':type(exc).__name__}]}
            # Run even when parsing port data failed; no second SSH session.
            location=localizar.collect(d,pw,secret,args,folder,connection=connection)
    except Exception as exc:
        error=mapear.ssh_diagnostic(exc,credentials)
        location['errors'].append({'command':'SSH','error':error})
        if not ports:ports={'Resumo':[{**meta,'status':'FALHA'}],'Ocorrencias':[{**meta,'comando':'SSH','erro':error}]}
    return ports,location

def add_sheet(book,title,headers,rows,widths):
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    sheet=book.create_sheet(title);sheet.append(headers)
    for row in rows:sheet.append([re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',str(v or ''))[:32767] for v in row])
    for cells in sheet:
        for cell in cells:cell.data_type='s';cell.alignment=Alignment(vertical='top',wrap_text=True)
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='203B60')
    for i,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(i)].width=width
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    return sheet

def export(tables,ports,path,quick):
    from modelo_excel import prepare_template
    from portas_excel import build_rows,vlan_rows
    book=prepare_template(tables)
    if not quick:
        rows,_=build_rows(ports)
        indexes=list(range(14))+[20]
        add_sheet(book,'Portas',['Switch','Porta','Descrição','Link','Admin','Velocidade','Duplex','Modo','VLAN acesso','VLANs trunk','VLAN nativa','Agregação','LACP papel','LACP estado','LLDP'],[[r[i] for i in indexes] for r in rows],[25,18,35,20,18,20,16,16,18,30,18,20,18,20,50])
        add_sheet(book,'VLANs',['Switch','VLAN ID','Nome','Instância','Descrição','Estado','Interfaces','Interface L3'],[r[:8] for r in vlan_rows(ports)],[25,15,28,25,35,18,45,25])
        pending=book['Pendências'];names=sorted({r['equipamento'] for r in ports.get('Ocorrencias',[])})
        for name in names:
            pending.append([name,'Portas: consultas incompletas','Conferir detalhes da coleta; portas já coletadas foram preservadas.'])
            for cell in pending[pending.max_row]:cell.data_type='s'
        pending.auto_filter.ref=pending.dimensions
        book.move_sheet(book['Portas'],offset=-2)
        book.move_sheet(book['VLANs'],offset=-1)
    book.active=0
    with path.open('xb') as f:book.save(f)

def run_collection(jobs,fw_jobs,targets,args,output_dir=None,on_event=None):
    if args.rapido and not targets:raise ValueError('Informe MACs para a consulta rápida.')
    run=Path(output_dir or localizar.BASE/'resultados')/datetime.now().strftime('%Y%m%d_%H%M%S_%f');logs=run/'logs';logs.mkdir(parents=True)
    fw_jobs=[] if args.rapido else fw_jobs
    total=len(jobs)+len(fw_jobs);completed=0
    def notify(message):
        if on_event:on_event({'message':message,'completed':completed,'total':total})
    ports={};locations=[];firewalls=[]
    for d,pw in fw_jobs:
        notify('Consultando firewall '+d['nome'])
        try:result=fortigate.collect(d,pw,args,logs)
        except Exception as exc:result={'name':d['nome'],'host':d['host'],'at':'','hosts':[],'system':{},'status':{},'ha':'','errors':[{'command':'SSH','error':type(exc).__name__}]}
        firewalls.append(result);completed+=1;notify(d['nome']+': consulta concluída')
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures={pool.submit(collect_switch,d,pw,secret,args,logs):d for d,pw,secret in jobs}
        for future in as_completed(futures):
            d=futures[future]
            try:p,l=future.result()
            except Exception as exc:
                l=failed(d);l['errors']=[{'command':'Coletor','error':type(exc).__name__}]
                p={'Ocorrencias':[{'equipamento':d['nome'],'comando':'Coletor','erro':type(exc).__name__}]}
            for key,rows in p.items():ports.setdefault(key,[]).extend(rows)
            locations.append(l);completed+=1;notify(d['nome']+': '+('consultado' if l['fdb_ok'] else 'tabela MAC indisponível'))
            for e in l['errors']:notify(d['nome']+' / '+e['command']+': '+e['error'])
            for e in p.get('Ocorrencias',[]):notify(d['nome']+' / '+e['comando']+': '+e['erro'])
    locations.sort(key=lambda d:d['name'])
    selected=targets if args.rapido else fortigate.discover(firewalls,locations,targets)
    tables=localizar.analyze(selected,locations);fortigate.add_tables(tables,firewalls)
    path=run/('localizacao_macs.xlsx' if args.rapido else 'mapeamento_rede.xlsx')
    export(tables,ports,path,args.rapido)
    (run/'coleta.json').write_text(json.dumps({'switches':locations,'firewalls':firewalls,'portas':ports},ensure_ascii=False,indent=2),encoding='utf-8')
    code=2 if ports.get('Ocorrencias') or any(d['errors'] for d in locations+firewalls) else 0
    from refresh_core.storage import event
    event('Mapeamento de Rede',path,status='Parcial' if code else 'Concluído',detail=str(len(jobs))+' switches')
    notify('Excel salvo: '+str(path))
    return {'path':path,'tables':tables,'code':code,'portas':ports}
