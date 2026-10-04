"""Auditoria somente leitura, com evidências selecionadas e falhas explícitas."""
import csv
import json
import re
from pathlib import Path
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from operacional import parse
from fortigate_ssh import SSH, COMMANDS

COMMANDS.update({'BGP':'get router info bgp summary','VPN':'get vpn ipsec tunnel summary',
 'Interfaces_cfg':'show full-configuration system interface','Global':'show full-configuration system global'})
OPERATIONAL_COMMANDS={key:COMMANDS[key] for key in ('Sistema','HA','BGP','VPN','Interfaces_cfg','Global')}
DEFAULTS={'hostname':True,'modelo':True,'versao':True,'ha':True,'bgp':True,'vpn':True,
          'acesso_inseguro':True,'timeout_admin':True,'timeout_maximo_minutos':10}
FIELDS=['nome','host','porta','usuario','modelo_esperado','versao_esperada','hostname_esperado','ha_esperado','bgp_esperados','vpns_esperadas']
UNKNOWN='Não foi possível verificar'

def policy(value):
    if set(value)!=set(DEFAULTS): raise ValueError('As regras devem conter exatamente os campos do modelo.')
    for key in DEFAULTS:
        if key!='timeout_maximo_minutos' and type(value[key]) is not bool: raise ValueError(key+': use true ou false')
    n=value['timeout_maximo_minutos']
    if type(n) is not int or not 1<=n<=480: raise ValueError('Timeout deve estar entre 1 e 480 minutos.')
    return dict(value)

def inventory(path):
    text=Path(path).read_text(encoding='utf-8-sig')
    if not text.strip(): raise ValueError('Inventário vazio')
    import io
    reader=csv.DictReader(io.StringIO(text),delimiter=';' if ';' in text.splitlines()[0] else ',')
    if not {'nome','host'}<=set(reader.fieldnames or []): raise ValueError('CSV precisa de nome e host.')
    return validate_devices(list(reader))

def validate_devices(items,allow_empty=False):
    rows=[]; seen=set()
    for number,row in enumerate(items,2):
        d={k:str(row.get(k) or '').strip() for k in FIELDS}
        if not d['nome'] or not d['host']: raise ValueError(f'Linha {number}: nome/host obrigatório')
        if any(c.isspace() for c in d['host']) or any(c in d['host'] for c in '/\\@?#'):raise ValueError('Host deve ser IP ou nome DNS, sem protocolo, porta ou espaços')
        if any(ord(c)<32 for v in d.values() for c in v):raise ValueError('Campos não podem conter caracteres de controle')
        d['porta']=int(d['porta'] or 22)
        if not 1<=d['porta']<=65535: raise ValueError('Porta SSH inválida')
        endpoint=(d['host'].lower(),d['porta'])
        if endpoint in seen: raise ValueError(f'Equipamento repetido: {d["host"]}')
        seen.add(endpoint)
        if d['ha_esperado'] not in ('','standalone','a-p','a-a'): raise ValueError('HA esperado: standalone, a-p ou a-a')
        rows.append(d)
    if not rows and not allow_empty: raise ValueError('Inventário vazio')
    return rows

def evaluate(device,data,rules):
    findings=[]
    def add(key,category,check):
        if not rules[key]: return
        response=data.get(category,{})
        if response.get('status')!='ok': state,evidence=UNKNOWN,response.get('error','Consulta indisponível')
        else:
            try: state,evidence=check(response['rows'])
            except (KeyError,ValueError,TypeError): state,evidence=UNKNOWN,'Campos necessários ausentes ou formato não reconhecido'
        findings.append([device['nome'],device['host'],key,state,evidence])
    def verdict(ok,evidence): return ('Conforme' if ok else 'Não conforme',evidence)
    def expected(rows,key,label):
        wanted=device.get('versao_esperada' if key=='versao' else key+'_esperado','')
        if not wanted: return 'Não aplicável','Valor esperado não definido no CSV'
        values={r['id']:r['value'] for r in rows}
        actual=values[label]
        if key=='modelo': actual=actual.split()[0]
        if key=='versao': actual=re.search(r'\bv(\d+\.\d+\.\d+)\b',actual)[1]
        return verdict(actual.casefold()==wanted.casefold(),f'Esperado: {wanted}; observado: {actual}')
    for key,label in [('hostname','Hostname'),('modelo','Version'),('versao','Version')]:
        add(key,'Sistema',lambda rows,k=key,l=label:expected(rows,k,l))
    def ha(rows):
        wanted=device.get('ha_esperado','')
        if not wanted: return 'Não aplicável','Modo HA esperado não definido'
        system={r['id']:r['value'] for r in data['Sistema']['rows']}
        mode=system['Current HA mode'].lower()
        if wanted=='standalone': return verdict('standalone' in mode,'Modo: '+mode)
        if wanted not in mode: return verdict(False,'Modo observado: '+mode+'; esperado: '+wanted)
        values={r['id']:r['value'] for r in rows}
        members={k[7:] for k in values if k.startswith('member/')}
        if len(members)<2: return verdict(False,'Menos de dois membros HA observados')
        if any('sync/'+m not in values for m in members): return UNKNOWN,'Sincronismo de todos os membros não identificado'
        return verdict(values['HA Health Status'].lower()=='ok' and all(values['sync/'+m]=='in-sync' for m in members),'Saúde: '+values['HA Health Status']+'; '+', '.join(m+': '+values['sync/'+m] for m in sorted(members)))
    add('ha','Sistema' if device.get('ha_esperado')=='standalone' else 'HA',ha)
    def sessions(rows,kind):
        field='neighbor' if kind=='bgp' else 'name'
        expected_set=set(filter(None,(x.strip() for x in device.get('bgp_esperados' if kind=='bgp' else 'vpns_esperadas','').split('|'))))
        if not rows and not expected_set: return 'Não aplicável','Nenhuma sessão observada; lista esperada não definida'
        observed={r[field] for r in rows}; missing=expected_set-observed
        bad=[r[field] for r in rows if (r['state']!='Established' if kind=='bgp' else r['total']==0 or r['up']!=r['total'])]
        return verdict(not missing and not bad,f'Sessões observadas: {len(rows)}; ausentes: {", ".join(sorted(missing)) or "nenhuma"}; inativas/parciais: {", ".join(bad) or "nenhuma"}')
    add('bgp','BGP',lambda rows:sessions(rows,'bgp'))
    add('vpn','VPN',lambda rows:sessions(rows,'vpn'))
    def access(rows):
        if not rows or any('allowaccess' not in r for r in rows): return UNKNOWN,'allowaccess não identificado em todas as interfaces'
        bad=[r['id'] for r in rows if {'http','telnet'} & set(r['allowaccess'].split())]
        return verdict(not bad,'Interfaces com HTTP/Telnet: '+(', '.join(bad) or 'nenhuma'))
    add('acesso_inseguro','Interfaces_cfg',access)
    def timeout(rows):
        value=int(rows[0]['admintimeout'])
        return verdict(0<value<=rules['timeout_maximo_minutos'],f'Timeout administrativo: {value} min; máximo: {rules["timeout_maximo_minutos"]} min')
    add('timeout_admin','Global',timeout)
    return findings

def collect(device,password,timeout,known_hosts,rules):
    data={}; connection=None
    try:
        connection=SSH(device,password,timeout,known_hosts)
        for category,command in OPERATIONAL_COMMANDS.items():
            if category in ('ARP','DHCP'): continue
            text=connection.command(command)
            if re.search(r'(?i)(command fail|parse error|permission denied|unknown action|return code -)',text):
                data[category]={'status':'error','error':'Comando recusado ou sem permissão'}; continue
            if category=='Global':
                match=re.search(r'^\s*set admintimeout (\d+)\s*$',text,re.M)
                data[category]={'status':'ok' if match and text.rstrip().endswith('end') else 'unparsed','rows':[{'admintimeout':match[1]}] if match else []}
            else: data[category]=parse(category,text)
            if category=='Sistema':
                values={r['id']:r['value'] for r in data[category]['rows']}
                if values.get('Virtual domain configuration','').lower()!='disable':
                    data={}
                    raise ValueError('Escopo sem VDOM não confirmado')
    except Exception as exc:
        # Não persistir mensagens de bibliotecas que possam conter credenciais.
        reason={'AuthenticationException':'Autenticação SSH recusada','BadHostKeyException':'Chave SSH mudou','TimeoutError':'Tempo esgotado','SSHException':'Falha SSH; verifique chave conhecida e negociação','ValueError':'Escopo sem VDOM não confirmado'}.get(type(exc).__name__,'Falha de conexão/coleta: '+type(exc).__name__)
        for category in OPERATIONAL_COMMANDS:
            if category not in data: data[category]={'status':'error','error':reason}
    finally:
        if connection: connection.close()
    return evaluate(device,data,rules)

def export(findings,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook(); ws=wb.active; ws.title='Resultados'
    ws.append(['Equipamento','Host','Regra','Resultado','Evidência'])
    for row in findings:
        ws.append(row)
        for cell in ws[ws.max_row]: cell.data_type='s'
    summary=wb.create_sheet('Resumo'); summary.append(['Equipamento','Host','Conforme','Não conforme','Não aplicável',UNKNOWN])
    keys=sorted({(r[0],r[1]) for r in findings})
    for name,host in keys:
        summary.append([name,host]+[sum(r[0]==name and r[1]==host and r[3]==s for r in findings) for s in ['Conforme','Não conforme','Não aplicável',UNKNOWN]])
        summary.cell(summary.max_row,1).data_type='s'; summary.cell(summary.max_row,2).data_type='s'
    for sheet in wb:
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='203E60')
        for i in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(i)].width=80 if sheet==ws and i==5 else 30
        for row in sheet.iter_rows(min_row=2):
            for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
    wb.save(path)

def run(devices,password,rules,workers=4,timeout=45,known_hosts='',progress=lambda x:None):
    policy(rules)
    if not any(rules[k] for k in rules if k!='timeout_maximo_minutos'): raise ValueError('Habilite pelo menos uma regra')
    findings=[]
    with ThreadPoolExecutor(max_workers=workers) as pool:
        tasks={pool.submit(collect,d,password,timeout,known_hosts,rules):d for d in devices}
        for future in as_completed(tasks):
            device=tasks[future]
            try: findings.extend(future.result())
            except Exception: findings.extend(evaluate(device,{},rules))
            progress(device['nome'])
    folder=Path(__file__).parent/'resultados'; folder.mkdir(exist_ok=True)
    path=folder/('auditoria_'+datetime.now().strftime('%Y%m%d_%H%M%S_%f')+'.xlsx')
    export(sorted(findings),path)
    path.with_suffix('.regras.json').write_text(json.dumps(rules,indent=2,ensure_ascii=False),encoding='utf-8')
    from refresh_core.storage import event
    event('Relatório auditoria',path,detail=str(len(devices))+' equipamentos; '+str(sum(r[3]=='Não conforme' for r in findings))+' não conformidades')
    return path
