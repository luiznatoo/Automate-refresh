import sys,unittest,uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from collections import Counter
from openpyxl import load_workbook
sys.path.insert(0,str(Path('outputs/localizador-mac').resolve()))
import rede
class Connection:
 def __init__(self,fail_mac=False):self.calls=Counter();self.fail_mac=fail_mac
 def __enter__(self):return self
 def __exit__(self,*args):pass
 def enable(self):pass
 def send_command(self,cmd,**kwargs):
  self.calls[cmd]+=1
  if cmd=='show running-config':return 'hostname SW01\ninterface GigabitEthernet1/0/1\n description PC\n switchport mode access\n switchport access vlan 10\ninterface GigabitEthernet1/0/2\n switchport mode trunk\n'
  if cmd=='show mac address-table':
   if self.fail_mac:raise TimeoutError('test')
   return 'Mac Address Table\n10 aabb.ccdd.eeff DYNAMIC Gi1/0/1\n'
  if cmd=='show lldp neighbors detail':return ''
  return ''
class IntegrationTests(unittest.TestCase):
 def run_case(self,quick=False,fail_mac=False,fw=False):
  folder=Path('work/localizador-tests/artifacts')/uuid.uuid4().hex;folder.mkdir(parents=True)
  args=SimpleNamespace(rapido=quick,timeout=5,workers=2,known_hosts=None,modo='lista')
  d={'nome':'SW01','host':'192.0.2.1','porta':'22','usuario':'operator','plataforma':'cisco_ios'}
  conn=Connection(fail_mac)
  with patch.object(rede,'connect_switch',return_value=conn) as connect,patch('netmiko.utilities.get_structured_data_textfsm',return_value=[]),patch.object(rede.fortigate,'collect',side_effect=TimeoutError('no firewall')):
   result=rede.run_collection([(d,'test','')],[({'nome':'FW','host':'192.0.2.254'},'test')] if fw else [],[{'mac':'aabbccddeeff','nome':'PC'}] if quick else [],args,folder)
   self.assertEqual(connect.call_count,1)
  return result,conn
 def test_one_session_shared_queries_and_workbook(self):
  result,c=self.run_case();self.assertEqual(c.calls['show running-config'],1);self.assertEqual(c.calls['show lldp neighbors detail'],1)
  self.assertEqual(result['tables']['Localização'][0]['Dispositivo'],'Nome não identificado')
  book=load_workbook(result['path'])
  self.assertEqual(book.sheetnames,['Portas','Dispositivos','VLANs','Pendências']);self.assertEqual(book['Portas'].max_row,3);self.assertEqual(book['Dispositivos']['G2'].value,'Localizado');book.close()
 def test_mac_failure_preserves_ports(self):
  result,c=self.run_case(fail_mac=True)
  book=load_workbook(result['path']);self.assertEqual(book['Portas'].max_row,3);self.assertGreater(book['Pendências'].max_row,1);book.close()
 def test_quick_skips_port_mapping(self):
  with patch.object(rede.mapear,'collect',side_effect=AssertionError('must not collect ports')):
   result,c=self.run_case(quick=True)
  book=load_workbook(result['path']);self.assertNotIn('Portas',book.sheetnames);book.close();self.assertNotIn('show version',c.calls)
 def test_firewall_failure_preserves_switch_data(self):
  result,c=self.run_case(fw=True)
  self.assertTrue(result['portas']['Interfaces_config']);self.assertEqual(len(result['tables']['Localização']),1);self.assertEqual(result['code'],2)
class OutputRegressionTests(unittest.TestCase):
 def test_lldp_numeric_id_and_interface_consolidated(self):
  from portas_excel import lldp_lines
  rows=[{'neighbor_name':'SW02','neighbor_interface':'ge-0/0/3.0','chassis_id':'00:10:db:ff:20:00'},
        {'neighbor_name':'SW02','neighbor_interface':'529','port_description':'ge-0/0/3.0','chassis_id':'00:10:db:ff:20:00'}]
  self.assertEqual(lldp_lines(rows),['SW02 | ge-0/0/3.0 | 00:10:db:ff:20:00'])
  rows.append({**rows[0],'neighbor_interface':'ge-0/0/4.0'})
  self.assertEqual(len(lldp_lines(rows)),2)
 def test_restored_excel_columns(self):
  folder=Path('work/localizador-tests/artifacts')/uuid.uuid4().hex;folder.mkdir(parents=True)
  ports={'Interfaces_config':[{'equipamento':'SW','interface':'ge-0/0/1','modo_configurado':'access'}],
         'PoE':[{'equipamento':'SW','interface':'ge-0/0/1','status':'ON'}],
         'Detalhes_portas':[{'equipamento':'SW','interface':'ge-0/0/1','media_type':'Copper','input_errors':'0','output_errors':'2'}]}
  path=folder/'test.xlsx';rede.export({'Localização':[],'Ocorrências':[]},ports,path,False)
  book=load_workbook(path);sheet=book['Portas'];headers=[c.value for c in sheet[1]]
  self.assertEqual(sheet.cell(2,headers.index('PoE')+1).value,'ON');self.assertEqual(sheet.cell(2,headers.index('Mídia / transceptor')+1).value,'Copper')
  self.assertIn('STP estado',headers);self.assertIn('Erros RX / TX',headers);book.close()
 def test_results_tab_opens_and_replaces(self):
  import interface as ui
  root=ui.tk.Tk();root.withdraw()
  with patch.object(ui.App,'load_initial'),patch('refresh_core.inventory.preload'):app=ui.App(root)
  try:
   self.assertEqual(len(app.workspace.tabs()),1)
   result={'path':'test.xlsx','tables':{'Localização':[],'Ocorrências':[],'A revisar':[]},'portas':{'Resumo':[]},'code':0}
   app.show_result(result);self.assertEqual(len(app.workspace.tabs()),2);self.assertEqual(app.workspace.select(),str(app.result_page))
   self.assertIn('Portas',app.preview_tables);app.show_result(result);self.assertEqual(len(app.workspace.tabs()),2)
  finally:app.close()
if __name__=='__main__':unittest.main()
